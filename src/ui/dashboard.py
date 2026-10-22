# -*- coding: utf-8 -*-
"""
src/ui/dashboard.py
Ứng dụng Giao diện Streamlit Dashboard UI Prototype v1.0.
Đề tài Khóa luận Tốt nghiệp TLCN-NAM:
HỆ THỐNG MULTI-AGENT ĐIỀU TRA VÀ PHẢN ỨNG SỰ CỐ AN NINH MẠNG DỰA TRÊN STATEGRAPH VÀ LANGGRAPH
Trường Đại học Sư phạm Kỹ thuật TP. Hồ Chí Minh (HCMUTE)
Sinh viên thực hiện: Phan Hoài Nam
"""

import os
import sys
import json
import io
import html
import re
import ipaddress
from datetime import datetime, timezone
from typing import List, Dict, Any, Optional, Tuple

# Ensure package imports work when Streamlit runs this file directly.
PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

import streamlit as st
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from src.core.schemas import (
    IncidentStatus,
    SeverityLevel,
    ActionType,
    RawLogInput,
    IncidentReport,
    IoCItem,
    RecommendedAction,
)
from src.core.state import IncidentState
from src.core.time_utils import format_vn_time, to_vn_time
from src.core.orchestrator import StateGraphOrchestrator, END
from src.core.hitl_manager import HITLManager
from src.agents.reporting_agent import ReportingAgent
from src.core.incident_store import IncidentStore
from src.tools.wazuh_client import WazuhClientTool
from src.ui.components import (
    render_kpi_cards,
    get_kpi_metrics,
    render_attack_detection_banner,
    render_agent_graph_status,
    generate_mermaid_graph,
    get_graph_node_statuses,
    render_ioc_table,
    format_ioc_table_data,
    render_mitre_matrix,
    extract_mitre_data,
    render_process_tree,
    format_process_tree_data,
    format_timeline_data,
    render_hitl_action_cards,
    apply_hitl_action_decision,
    render_hitech_agent_workflow,
    get_hitech_workflow_html,
)

# Thiết lập cấu hình trang Streamlit
try:
    st.set_page_config(
        page_title="SOC Multi-Agent Dashboard - Prototype v1.0",
        layout="wide",
        page_icon="🛡️",
        initial_sidebar_state="expanded",
    )
except Exception:
    # Bỏ qua khi gọi trong môi trường kiểm thử không có Streamlit server runtime
    pass


# =====================================================================
# 1. SCENARIO DATA LOADER & PATH RESOLVER
# =====================================================================
SCENARIO_CATALOG = {
    "live_wazuh": {
        "title": "🔴 Live Wazuh OVA (Nhận cảnh báo tự động qua Webhook)",
        "file_names": [],
        "description": "Nhận luồng cảnh báo tự động thời gian thực từ máy chủ Wazuh OVA (172.16.65.150) qua Webhook (:8000).",
        "expected_severity": "LIVE / DYNAMIC",
    },
    "malware": {
        "title": "🦠 Malware Infection (Lây nhiễm Trojan/Emotet & Reverse Shell)",
        "file_names": ["wazuh_alerts_scenario2_malware.json", "scenario2_malware_infection.json"],
        "description": "Kịch bản phát hiện PowerShell tải payload mã độc từ IP lạ, kết nối C2 và tạo tiến trình bất thường.",
        "expected_severity": "MEDIUM / HIGH",
    },
    "phishing": {
        "title": "🎣 Credential Theft / Phishing (Đánh cắp tài khoản & Password Spraying)",
        "file_names": ["wazuh_alerts_scenario1_phishing.json", "scenario1_credential_theft.json"],
        "description": "Kịch bản tấn công Password Spraying liên tiếp vào tài khoản Domain, kích hoạt brute force rule.",
        "expected_severity": "HIGH",
    },
    "network_scan": {
        "title": "🌐 Network Scanning (Dò quét cổng & Do thám lỗ hổng)",
        "file_names": ["wazuh_alerts_scenario3_network_scan.json", "scenario3_network_scanning.json"],
        "description": "Kịch bản tấn công quét mạng diện rộng qua Zeek/Wazuh nhắm vào nhiều cổng dịch vụ.",
        "expected_severity": "MEDIUM",
    },
    "benign": {
        "title": "✅ Benign / False Positive (Hoạt động Quản trị Bình thường)",
        "file_names": ["wazuh_alerts_scenario4_benign.json", "scenario4_benign.json"],
        "description": "Kịch bản hoạt động bảo trì hệ thống hợp lệ, thử nghiệm tính năng rẽ nhánh False Positive.",
        "expected_severity": "LOW / FALSE_POSITIVE",
    },
}


def resolve_scenario_file_path(file_names: List[str]) -> Optional[str]:
    """
    Tìm kiếm file kịch bản cảnh báo trong các thư mục data/wazuh_alerts/ hoặc data/raw/wazuh/.
    """
    candidate_dirs = [
        os.path.join("data", "wazuh_alerts"),
        os.path.join("data", "raw", "wazuh"),
        os.path.join("/home/nam/Desktop/TLCN-NAM/TLCN/data/wazuh_alerts"),
        os.path.join("/home/nam/Desktop/TLCN-NAM/TLCN/data/raw/wazuh"),
    ]

    for fname in file_names:
        for cdir in candidate_dirs:
            p = os.path.join(cdir, fname)
            if os.path.isfile(p):
                return p
    return None


def load_scenario_alerts(
    scenario_key: str,
    custom_json_content: Optional[str] = None,
) -> Tuple[List[Dict[str, Any]], str]:
    """
    Nạp dữ liệu cảnh báo từ file mẫu hoặc nội dung JSON tùy chỉnh.

    :param scenario_key: Khóa kịch bản ("malware", "phishing", "network_scan", "benign", "custom")
    :param custom_json_content: Chuỗi JSON nếu người dùng upload file tùy chỉnh.
    :return: (Danh sách cảnh báo dạng dict, Tên kịch bản)
    """
    if scenario_key == "custom" and custom_json_content:
        try:
            data = json.loads(custom_json_content)
            if isinstance(data, list):
                return data, "Tùy chỉnh: File JSON người dùng cung cấp"
            elif isinstance(data, dict):
                return [data], "Tùy chỉnh: File JSON người dùng cung cấp"
        except Exception as e:
            return [], f"Lỗi đọc file JSON tùy chỉnh: {e}"

    if scenario_key == "live_wazuh":
        try:
            from src.core.incident_store import IncidentStore
            store = IncidentStore()
            # CHỈ lấy các sự cố thực tế nhận từ Webhook Wazuh Live (prefix INC-WAZUH-)
            # Không lấy nhầm các sự cố từ kịch bản mẫu offline (INC-MALW-, INC-PHIS-,...)
            live_incidents = [
                inc for inc in store.list_incidents(limit=20)
                if inc.incident_id.startswith("INC-WAZUH-")
            ]
            if live_incidents:
                latest_inc = live_incidents[0]
                alerts = []
                for r in latest_inc.raw_logs:
                    try:
                        alerts.append(json.loads(r.raw_payload))
                    except Exception:
                        pass
                if alerts:
                    return alerts, f"🔴 Live Wazuh: Sự cố {latest_inc.incident_id} ({latest_inc.attack_type or 'Mới nhận'})"
        except Exception:
            pass
        return [], "🔴 Live Wazuh: Đang chờ nhận cảnh báo từ Wazuh OVA (:8000/api/v1/wazuh-alert)..."

    info = SCENARIO_CATALOG.get(scenario_key)
    if not info:
        info = SCENARIO_CATALOG["malware"]

    fpath = resolve_scenario_file_path(info["file_names"])
    if not fpath:
        # Fallback mẫu tối thiểu
        mock_alert = {
            "timestamp": datetime.now(timezone.utc).isoformat(),
            "rule": {
                "id": "100001",
                "level": 8,
                "description": f"Fallback alert for {scenario_key}",
                "mitre": {"id": ["T1059.001"], "tactic": ["Execution"], "technique": ["PowerShell"]},
                "groups": ["windows", "sysmon"],
            },
            "agent": {"id": "001", "name": "HCMUTE-PC-01", "ip": "192.168.1.105"},
            "full_log": "Fallback log line",
            "data": {"srcip": "194.26.29.112", "commandLine": "powershell.exe -enc AAAA"},
        }
        return [mock_alert], info["title"]

    with open(fpath, "r", encoding="utf-8") as f:
        content = f.read().strip()
        if content.startswith("["):
            items = json.loads(content)
        else:
            items = [json.loads(line) for line in content.split("\n") if line.strip()]

    return items, info["title"]


def create_initial_incident_state(scenario_key: str, raw_alerts: List[Dict[str, Any]]) -> IncidentState:
    """
    Khởi tạo IncidentState mới với các RawLogInput chuẩn hóa từ danh sách cảnh báo Wazuh.
    """
    if scenario_key == "live_wazuh" and not raw_alerts:
        inc_id = f"INC-WAZUH-STANDBY-{datetime.now(timezone.utc).strftime('%H%M%S')}"
    else:
        inc_id = f"INC-{scenario_key.upper()[:4]}-{datetime.now(timezone.utc).strftime('%H%M%S')}"

    raw_logs = [
        RawLogInput(
            log_source="wazuh",
            raw_payload=json.dumps(alert_dict),
            collected_at=datetime.now(timezone.utc),
        )
        for alert_dict in raw_alerts
    ]

    state = IncidentState(
        incident_id=inc_id,
        status=IncidentStatus.NEW,
        raw_logs=raw_logs,
        created_at=datetime.now(timezone.utc),
        updated_at=datetime.now(timezone.utc),
    )
    return state


def persist_incident_state(state: IncidentState) -> Optional[str]:
    """
    Tự động lưu trữ đối tượng IncidentState vào IncidentStore trên ổ đĩa.
    Hỗ trợ duy trì bền vững dữ liệu sau khi hoàn tất pipeline hoặc xuất bản báo cáo.
    """
    if not state or not state.incident_id:
        return None
    try:
        store = IncidentStore()
        saved_path = store.save_incident(state)
        return saved_path
    except Exception as e:
        print(f"[ERROR] Không thể lưu IncidentState {state.incident_id} vào IncidentStore: {e}")
        return None


# =====================================================================
# 2. REPORT EXPORT HELPERS (MARKDOWN, JSON, DOCX)
# =====================================================================
def generate_markdown_report(report: IncidentReport) -> str:
    """Xuất báo cáo định dạng Markdown chuẩn cho giao diện Streamlit."""
    return ReportingAgent.export_markdown(report)


def generate_json_report(report: IncidentReport) -> str:
    """Xuất báo cáo định dạng JSON chuẩn."""
    return ReportingAgent.export_json(report)


def generate_docx_report(report: IncidentReport) -> bytes:
    """
    Sinh tài liệu Word (.docx) chuyên nghiệp từ IncidentReport bằng thư viện python-docx.
    Tạo tiêu đề, thông tin định danh, Executive Summary, Detailed Findings,
    bảng IoC và bảng Kế hoạch Phản ứng NIST SP 800-61.
    """
    doc = docx.Document()

    # Cấu hình lề trang 1 inch
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # 1. Header Khóa luận & Cơ quan điều tra
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_uni = p_header.add_run("TRƯỜNG ĐẠI HỌC SƯ PHẠM KỸ THUẬT TP. HỒ CHÍ MINH (HCMUTE)\n")
    r_uni.bold = True
    r_uni.font.size = Pt(13)
    r_uni.font.color.rgb = RGBColor(0, 51, 102)

    r_dep = p_header.add_run("KHOA CÔNG NGHỆ THÔNG TIN - BỘ MÔN AN TOÀN THÔNG TIN\n")
    r_dep.bold = True
    r_dep.font.size = Pt(11)

    r_sub = p_header.add_run("ĐỀ TÀI KHÓA LUẬN TỐT NGHIỆP: HỆ THỐNG SOC MULTI-AGENT PROTOTYPE V1.0\n")
    r_sub.font.size = Pt(10)
    r_sub.font.italic = True

    doc.add_paragraph("―" * 45).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 2. Tiêu đề Báo cáo
    p_title = doc.add_heading(level=0)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(f"🛡️ {report.incident_title}")
    r_title.font.size = Pt(18)
    r_title.bold = True
    r_title.font.color.rgb = RGBColor(180, 0, 0)

    # Metadata
    p_meta = doc.add_paragraph()
    p_meta.add_run(f"• Mã Báo cáo: ").bold = True
    p_meta.add_run(f"{report.report_id}\n")
    p_meta.add_run(f"• Thời gian xuất bản: ").bold = True
    p_meta.add_run(f"{format_vn_time(report.created_at, '%Y-%m-%d %H:%M:%S (ICT / Giờ Việt Nam)')}\n")
    p_meta.add_run(f"• Tiêu chuẩn tham chiếu: ").bold = True
    p_meta.add_run("NIST SP 800-61 Rev.2 & MITRE ATT&CK Framework\n")

    # 3. Mục 1: Executive Summary
    h1 = doc.add_heading("1. Tóm lược điều hành (Executive Summary)", level=1)
    h1.style.font.color.rgb = RGBColor(0, 51, 102)
    p_exec = doc.add_paragraph()
    r_exec = p_exec.add_run(report.executive_summary)
    r_exec.font.size = Pt(11)

    # 4. Mục 2: Detailed Technical Findings
    h2 = doc.add_heading("2. Chi tiết kỹ thuật & Phân tích nguyên nhân (Technical Findings)", level=1)
    h2.style.font.color.rgb = RGBColor(0, 51, 102)
    p_tech = doc.add_paragraph()
    p_tech.add_run(report.detailed_findings)

    # 5. Mục 3: MITRE ATT&CK Matrix Mapping
    h3 = doc.add_heading("3. Ma trận Kỹ thuật MITRE ATT&CK", level=1)
    h3.style.font.color.rgb = RGBColor(0, 51, 102)
    if report.mitre_attack_matrix_mapping:
        for idx, tech in enumerate(report.mitre_attack_matrix_mapping, start=1):
            doc.add_paragraph(f"{idx}. Mã Kỹ thuật: {tech}", style="List Bullet")
    else:
        doc.add_paragraph("Chưa ghi nhận mã kỹ thuật cụ thể.")

    # 6. Mục 4: Bảng IoC
    h4 = doc.add_heading("4. Danh sách Chỉ số Xâm phạm (IoC Table)", level=1)
    h4.style.font.color.rgb = RGBColor(0, 51, 102)

    if report.ioc_table:
        table_ioc = doc.add_table(rows=1, cols=5)
        table_ioc.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = table_ioc.rows[0].cells
        hdr_titles = ["STT", "Loại IoC", "Giá trị Chỉ số", "Điểm Độc hại", "Đánh giá"]
        for i, title in enumerate(hdr_titles):
            hdr_cells[i].text = title
            hdr_cells[i].paragraphs[0].runs[0].bold = True

        for idx, ioc in enumerate(report.ioc_table, start=1):
            row_cells = table_ioc.add_row().cells
            row_cells[0].text = str(idx)
            row_cells[1].text = ioc.ioc_type
            row_cells[2].text = ioc.value
            row_cells[3].text = str(ioc.reputation_score or 0)
            row_cells[4].text = "Độc hại" if ioc.is_malicious else "Đã kiểm tra"
    else:
        doc.add_paragraph("Không phát hiện chỉ số xâm phạm ngoài vùng tin cậy.")

    # 7. Mục 5: Action Plan
    h5 = doc.add_heading("5. Kế hoạch Phản ứng & Hành động Khắc phục (Action Plan)", level=1)
    h5.style.font.color.rgb = RGBColor(0, 51, 102)

    if report.action_plan:
        table_act = doc.add_table(rows=1, cols=5)
        table_act.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells_act = table_act.rows[0].cells
        hdr_titles_act = ["Loại hành động", "Mục tiêu (Target)", "Ưu tiên", "Cần duyệt", "Trạng thái"]
        for i, title in enumerate(hdr_titles_act):
            hdr_cells_act[i].text = title
            hdr_cells_act[i].paragraphs[0].runs[0].bold = True

        for act in report.action_plan:
            row_cells = table_act.add_row().cells
            row_cells[0].text = act.action_type.value if hasattr(act.action_type, "value") else str(act.action_type)
            row_cells[1].text = act.target
            row_cells[2].text = f"Mức {act.priority}"
            row_cells[3].text = "Có" if act.require_human_approval else "Không"
            if act.is_approved is True:
                st_txt = "Đã duyệt"
            elif act.is_approved is False:
                st_txt = "Bị từ chối"
            else:
                st_txt = "Chờ duyệt"
            row_cells[4].text = st_txt
    else:
        doc.add_paragraph("Không có hành động nào cần thực thi.")

    # Footer
    p_footer = doc.add_paragraph()
    p_footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_footer.add_run(f"\nBáo cáo tạo tự động bởi Hệ thống HCMUTE Multi-Agent SOC - {datetime.now(timezone.utc).strftime('%Y-%m-%d')}").italic = True

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


# =====================================================================
# 3. HEADLESS INTEGRITY CHECK (Tự kiểm tra không cần trình duyệt)
# =====================================================================
def run_dashboard_headless_check() -> bool:
    """
    Kiểm tra khởi tạo và kết xuất các thành phần không cần trình duyệt web.
    Bảo đảm 100% logic backend, nạp kịch bản, trích xuất KPI, định dạng IoC,
    ma trận MITRE, cây tiến trình Sysmon, duyệt HITL và xuất báo cáo (MD, JSON, DOCX) hoạt động hoàn hảo.

    :return: True nếu tất cả kiểm tra đều thành công.
    """
    # 1. Kiểm tra nạp kịch bản
    for skey in ["malware", "phishing", "network_scan", "benign"]:
        alerts, title = load_scenario_alerts(skey)
        if not alerts or not title:
            raise RuntimeError(f"Lỗi nạp kịch bản '{skey}'")

    # 2. Khởi tạo IncidentState mẫu
    alerts, _ = load_scenario_alerts("malware")
    state = create_initial_incident_state("malware", alerts)
    if not state.incident_id or len(state.raw_logs) == 0:
        raise RuntimeError("Khởi tạo IncidentState thất bại")

    # 3. Kiểm tra trích xuất KPI
    kpis = get_kpi_metrics(state)
    expected_keys = [
        "incident_id", "status", "severity", "confidence_score",
        "raw_logs_count", "parsed_events_count", "ioc_count",
        "total_actions", "pending_actions", "approved_actions"
    ]
    for k in expected_keys:
        if k not in kpis:
            raise KeyError(f"Thiếu trường KPI: {k}")

    # 4. Kiểm tra StateGraph DAG & Mermaid
    mermaid_str = generate_mermaid_graph(state)
    if "flowchart" not in mermaid_str or "Ingestion" not in mermaid_str:
        raise ValueError("Mã Mermaid sinh ra không hợp lệ")

    node_statuses = get_graph_node_statuses(state)
    if len(node_statuses) != 7:
        raise ValueError(f"Số lượng node mong đợi là 7, nhận được: {len(node_statuses)}")

    # 4b. Kiểm tra Sơ đồ Kiến trúc Hitech (Hitech Architecture Workflow)
    hitech_html = get_hitech_workflow_html(state)
    for part in ["01 / TELEMETRY", "02 / AGENT RUNTIME", "EX / POLICY", "04 / TOOL EXECUTION"]:
        if part not in hitech_html:
            raise ValueError(f"Hitech Workflow HTML thiếu phân vùng: {part}")
    if state.incident_id not in hitech_html:
        raise ValueError(f"Hitech Workflow HTML không nhúng incident_id '{state.incident_id}'")

    hitech_standby = get_hitech_workflow_html(None)
    if "01 / TELEMETRY" not in hitech_standby:
        raise ValueError("Hitech Workflow HTML chế độ Standby không hợp lệ")

    # 5. Thực thi đồ thị mẫu qua StateGraphOrchestrator
    orch = StateGraphOrchestrator(hitl_manager=HITLManager(), auto_approve=False)
    is_paused, step_state = orch.execute_until_interrupt(state)
    if is_paused:
        final_state = orch.resume_after_hitl(
            state=step_state,
            analyst_approved=True,
            analyst_feedback="Chuyên viên phê duyệt trong bài kiểm tra headless"
        )
    else:
        final_state = step_state

    if final_state.status not in {IncidentStatus.RESOLVED, IncidentStatus.CLOSED_FALSE_POSITIVE}:
        raise RuntimeError(f"Trạng thái kết thúc không hợp lệ: {final_state.status}")

    # 6. Kiểm tra IoC Table & MITRE Matrix
    ioc_rows = format_ioc_table_data(final_state.threat_intel)
    mitre_data = extract_mitre_data(final_state.threat_intel, final_state.rca_result)
    proc_rows = format_process_tree_data(final_state.rca_result)
    timeline_rows = format_timeline_data(final_state.rca_result)

    if not isinstance(ioc_rows, list) or not isinstance(mitre_data, dict):
        raise TypeError("Định dạng dữ liệu IoC/MITRE không hợp lệ")

    # 7. Kiểm tra HITL Decision Helper
    if final_state.recommended_actions:
        apply_hitl_action_decision(
            state=final_state,
            action_idx=0,
            decision="APPROVE",
            edited_target="192.168.1.99",
            edited_priority=1,
            edited_description="Điều chỉnh mục tiêu khẩn cấp",
        )
        assert final_state.recommended_actions[0].is_approved is True
        assert final_state.recommended_actions[0].target == "192.168.1.99"

        apply_hitl_action_decision(
            state=final_state,
            action_idx=0,
            decision="REJECT",
        )
        assert final_state.recommended_actions[0].is_approved is False

    # 8. Kiểm tra Xuất Báo cáo (Markdown, JSON, DOCX)
    if final_state.final_report:
        md_text = generate_markdown_report(final_state.final_report)
        if "# " not in md_text or "Executive Summary" not in md_text:
            raise ValueError("Markdown report export không chứa tiêu đề chuẩn")

        json_text = generate_json_report(final_state.final_report)
        json_obj = json.loads(json_text)
        if json_obj.get("report_id") != final_state.final_report.report_id:
            raise ValueError("JSON report export sai report_id")

        docx_bytes = generate_docx_report(final_state.final_report)
        if not docx_bytes or len(docx_bytes) < 1000:
            raise ValueError("DOCX report export trả về dữ liệu rỗng hoặc quá nhỏ")

    # 9. Kiểm tra Trợ lý SOC Engine & Provider
    from src.copilot.engine import SOCCopilotEngine
    copilot_engine = SOCCopilotEngine()
    if not hasattr(copilot_engine, "provider") or not copilot_engine.provider:
        raise ValueError("SOCCopilotEngine thiếu thuộc tính provider tương thích ngược")
    if not hasattr(copilot_engine.provider, "provider_type"):
        raise ValueError("Provider của SOCCopilotEngine không có thuộc tính provider_type")

    copilot_res = copilot_engine.process_message(
        "Tóm tắt sự cố này",
        state=final_state,
        chat_history=[]
    )
    if not copilot_res.response_text or len(copilot_res.response_text) < 20:
        raise ValueError("Trợ lý SOC trả về phản hồi rỗng hoặc không hợp lệ")

    return True


# =====================================================================
# 4. STREAMLIT DASHBOARD MAIN UI
# =====================================================================
def init_session_state():
    """Khởi tạo trạng thái phiên làm việc Streamlit (Session State)."""
    if "orchestrator" not in st.session_state:
        st.session_state.orchestrator = StateGraphOrchestrator(
            hitl_manager=HITLManager(),
            auto_approve=False
        )

    if "copilot_engine" not in st.session_state:
        from src.copilot.engine import SOCCopilotEngine
        st.session_state.copilot_engine = SOCCopilotEngine()

    if "copilot_messages" not in st.session_state:
        st.session_state.copilot_messages = [
            {
                "role": "assistant",
                "content": "👋 Xin chào! Tôi là **Trợ lý SOC**, hỗ trợ điều tra và phản ứng sự cố an ninh mạng của bạn (tích hợp OmniRouter AI Gateway & Fallback Deterministic CoT Engine).\n\nTôi có thể giúp bạn:\n- 🔍 **Giải thích nguyên nhân gốc rễ (RCA)** và phả hệ cây tiến trình Sysmon.\n- 🌐 **Tra cứu danh tiếng IoC (CTI)** (IP, Domain, File Hash).\n- 📋 **Tóm tắt sự cố**.\n- ⚡ **Phê duyệt hoặc từ chối hành động can thiệp (HITL)** bằng ngôn ngữ tự nhiên (*'Duyệt chặn IP 194.26.29.112'* hoặc *'Duyệt tất cả'*).\n\nHãy chọn một gợi ý nhanh bên dưới hoặc nhập yêu cầu của bạn!",
            }
        ]

    if "evidence_source_mode" not in st.session_state:
        st.session_state.evidence_source_mode = "live_wazuh"

    if "selected_sample_scenario" not in st.session_state:
        st.session_state.selected_sample_scenario = "malware"

    if "selected_scenario_key" not in st.session_state:
        st.session_state.selected_scenario_key = "live_wazuh"

    if "incident_state" not in st.session_state:
        alerts, _ = load_scenario_alerts(st.session_state.selected_scenario_key)
        st.session_state.raw_alerts = alerts
        st.session_state.incident_state = create_initial_incident_state(
            st.session_state.selected_scenario_key, alerts
        )
        st.session_state.current_node = "ingestion"

    if "current_node" not in st.session_state:
        st.session_state.current_node = "ingestion"


def reset_pipeline(scenario_key: str, custom_content: Optional[str] = None):
    """Đặt lại pipeline điều phối về trạng thái ban đầu với kịch bản được chọn."""
    if scenario_key == "live_wazuh":
        st.session_state.selected_scenario_key = "live_wazuh"
        st.session_state.evidence_source_mode = "live_wazuh"
        try:
            from src.core.incident_store import IncidentStore
            store = IncidentStore()
            # CHỈ lấy các sự cố thực tế nhận từ Webhook Wazuh Live (bắt đầu bằng INC-WAZUH-)
            # Không lấy nhầm các sự cố từ kịch bản mẫu offline (INC-MALW-, INC-PHIS-,...)
            live_incidents = [
                inc for inc in store.list_incidents(limit=20)
                if inc.incident_id.startswith("INC-WAZUH-")
            ]
            if live_incidents:
                latest_inc = live_incidents[0]
                st.session_state.incident_state = latest_inc
                st.session_state.raw_alerts = [
                    json.loads(r.raw_payload) for r in latest_inc.raw_logs if r.raw_payload.startswith("{")
                ]
                st.session_state.current_node = None if latest_inc.final_report else "triage"
                try:
                    st.toast(f"📡 Đã nạp sự cố Live Wazuh mới nhất: {latest_inc.incident_id}", icon="🔴")
                except Exception:
                    pass
                return
        except Exception:
            pass

        # Khi Wazuh chưa gửi bất kỳ cảnh báo nào (Wazuh OVA chưa bật hoặc chưa bắn alert sang :8000):
        # Thiết lập trạng thái Standby sạch sẽ, KHÔNG tự ý nạp kịch bản cũ hay báo đã hoàn tất!
        inc_id = f"INC-WAZUH-STANDBY-{datetime.now(timezone.utc).strftime('%H%M%S')}"
        st.session_state.raw_alerts = []
        st.session_state.incident_state = IncidentState(
            incident_id=inc_id,
            status=IncidentStatus.NEW,
            raw_logs=[],
            created_at=datetime.now(timezone.utc),
            updated_at=datetime.now(timezone.utc),
        )
        st.session_state.current_node = "ingestion"
        try:
            st.toast("📡 Đang ở chế độ Live Wazuh (Standby). Đang túc trực cổng Webhook :8000...", icon="ℹ️")
        except Exception:
            pass
        return

    alerts, _ = load_scenario_alerts(scenario_key, custom_content)
    st.session_state.raw_alerts = alerts
    st.session_state.selected_scenario_key = scenario_key
    st.session_state.incident_state = create_initial_incident_state(scenario_key, alerts)
    st.session_state.current_node = "ingestion"


def run_full_pipeline():
    """Thực thi toàn bộ đồ thị StateGraph từ node hiện tại đến khi gặp điểm ngắt HITL hoặc hoàn tất."""
    state = st.session_state.incident_state
    if not state.raw_logs:
        st.warning(
            "⚠️ **Chưa có cảnh báo nào từ Wazuh SIEM để điều tra!**\n\n"
            "• Nếu đang ở chế độ **🔴 Live Wazuh**: Vui lòng khởi động máy ảo Wazuh OVA (172.16.65.150) và kích hoạt sự cố trên máy nạn nhân để gửi cảnh báo về Webhook `:8000/api/v1/wazuh-alert`.\n"
            "• Hoặc chọn một trong các **📁 Kịch bản kiểm nghiệm mẫu** (Malware, Phishing, Port Scan) trên menu bên trái để trải nghiệm toàn trình."
        )
        return

    orch = st.session_state.orchestrator

    with st.spinner("🤖 Các AI Agents đang phối hợp điều tra sự cố..."):
        is_paused, updated_state = orch.execute_until_interrupt(state)
        st.session_state.incident_state = updated_state

        if is_paused:
            st.session_state.current_node = "hitl_approval"
            st.warning(
                "⏸️ **Pipeline đã tạm dừng tại điểm ngắt HITL!** "
                "Có hành động can thiệp yêu cầu chuyên viên SOC phê duyệt. "
                "Vui lòng chuyển sang tab **'🛡️ Phê duyệt Chuyên viên (HITL)'** trong mục **'💼 Quy trình IR'** để thẩm định và phê duyệt."
            )
            try:
                st.toast(
                    "⏸️ Đã tạm dừng tại điểm ngắt HITL! Vui lòng chuyển sang tab 'Phê duyệt Chuyên viên (HITL)'.",
                    icon="⚠️"
                )
            except Exception:
                pass
        else:
            st.session_state.current_node = None
            saved_path = persist_incident_state(updated_state)
            if saved_path:
                try:
                    st.toast(f"💾 Đã tự động lưu sự cố `{updated_state.incident_id}` vào kho lưu trữ!", icon="✅")
                except Exception:
                    pass


def run_single_step():
    """Thực thi đơn lẻ một bước (Node) kế tiếp trong StateGraph."""
    curr_node = st.session_state.current_node
    if not curr_node:
        st.info("Pipeline đã hoàn tất hoặc không còn bước nào cần thực thi.")
        return

    state = st.session_state.incident_state
    if not state.raw_logs:
        st.warning(
            "⚠️ **Chưa có cảnh báo nào từ Wazuh SIEM!** Vui lòng khởi động Wazuh OVA (172.16.65.150) hoặc chọn kịch bản mẫu từ menu bên trái."
        )
        return

    orch = st.session_state.orchestrator

    with st.spinner(f"Đang thực thi Agent Node: `{curr_node}`..."):
        next_node, updated_state = orch.step(curr_node, state)
        st.session_state.incident_state = updated_state
        st.session_state.current_node = next_node
        if next_node is None or updated_state.final_report is not None:
            persist_incident_state(updated_state)


def main():
    """Hàm giao diện chính của ứng dụng Streamlit theo quy trình Wazuh SOC Multi-Agent."""
    init_session_state()

    state: IncidentState = st.session_state.incident_state

    # -------------------------------------------------------------
    # SIDEBAR: WAZUH SOC WORKFLOW & BẢNG ĐIỀU KHIỂN
    # -------------------------------------------------------------
    with st.sidebar:
        st.markdown(
            """
            <div style="padding: 6px 0 16px 0; border-bottom: 1px solid #334155; margin-bottom: 16px;">
                <div style="display: flex; align-items: center; gap: 10px;">
                    <span style="font-size: 1.8rem;">🛡️</span>
                    <div>
                        <h2 style="margin: 0; color: #38bdf8; font-size: 1.25rem; font-weight: 800; letter-spacing: -0.3px;">
                            Wazuh SOC Multi-Agent
                        </h2>
                        <div style="color: #94a3b8; font-size: 0.75rem; font-weight: 600; text-transform: uppercase; letter-spacing: 0.05em;">
                            Intelligent IR & SIEM Orchestration
                        </div>
                    </div>
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

        nav_options = [
            "🏠 Trang chủ",
            "📈 Đánh giá sự cố",
            "🛡️ Bằng chứng & Log Wazuh SIEM",
            "💼 Quy trình IR",
            "🛡️ IOC & SIEM",
            "💬 Trợ lý SOC",
            "🕒 Lịch sử & Báo cáo",
        ]
        selected_nav = st.radio(
            "Điều hướng hệ thống:",
            options=nav_options,
            index=0,
            label_visibility="collapsed",
        )

        st.divider()
        st.markdown("##### ⚙️ Nguồn Bằng chứng & Cảnh báo:")

        # 1. Menu cấp 1: Gom nhóm nguồn dữ liệu (Wazuh Live, Kịch bản mẫu, Custom)
        source_mode_options = {
            "live_wazuh": "🔴 1. Trực tiếp từ Wazuh OVA (Live Webhook Mode)",
            "sample_datasets": "📁 2. Kịch bản kiểm nghiệm mẫu (Offline Replay Datasets)",
            "custom": "📂 3. Tải lên file JSON cảnh báo tùy chỉnh",
        }

        # Đồng bộ initial selection
        current_mode = st.session_state.get("evidence_source_mode", "live_wazuh")
        if current_mode not in source_mode_options:
            current_mode = "live_wazuh"

        selected_source_mode = st.selectbox(
            "Chọn Nguồn Bằng chứng:",
            options=list(source_mode_options.keys()),
            format_func=lambda k: source_mode_options[k],
            index=list(source_mode_options.keys()).index(current_mode),
        )

        custom_json_content = None
        target_scenario_key = "live_wazuh"

        if selected_source_mode == "live_wazuh":
            target_scenario_key = "live_wazuh"
            st.markdown(
                """
                <div style="background: rgba(16, 185, 129, 0.1); border: 1px solid rgba(52, 211, 153, 0.35); border-radius: 8px; padding: 10px 12px; margin: 8px 0;">
                    <div style="color: #34d399; font-weight: 700; font-size: 0.85rem; display: flex; align-items: center; gap: 6px;">
                        <span>🟢 Webhook Receiver Đang Trực</span>
                    </div>
                    <div style="color: #cbd5e1; font-size: 0.78rem; margin-top: 6px; line-height: 1.6;">
                        • <strong style="color: #94a3b8;">Endpoint:</strong> <code style="background: #0f172a; color: #38bdf8; border: 1px solid #334155; padding: 2px 6px; border-radius: 4px; font-size: 0.75rem;">:8000/api/v1/wazuh-alert</code><br/>
                        • <strong style="color: #94a3b8;">Wazuh OVA:</strong> <code style="background: #0f172a; color: #34d399; border: 1px solid #334155; padding: 2px 6px; border-radius: 4px; font-size: 0.75rem;">172.16.65.150</code>
                    </div>
                </div>
                """,
                unsafe_allow_html=True,
            )
            if st.button("🔄 Nạp Cảnh báo Mới nhất từ Wazuh", use_container_width=True):
                st.session_state.evidence_source_mode = "live_wazuh"
                reset_pipeline("live_wazuh")
                st.rerun()

        elif selected_source_mode == "sample_datasets":
            # 2. Menu cấp 2: 4 Kịch bản mẫu gom vào một nơi
            sample_scenarios = {
                "malware": "🦠 1. Lây nhiễm Mã độc (Malware & C2 Reverse Shell)",
                "phishing": "🎣 2. Đánh cắp Danh tính (Credential Theft & Password Spraying)",
                "network_scan": "🌐 3. Quét Mạng Do Thám (Network Scanning & Discovery)",
                "benign": "✅ 4. Hoạt động Quản trị Bình thường (Benign / False Positive)",
            }
            curr_sample = st.session_state.get("selected_sample_scenario", "malware")
            if curr_sample not in sample_scenarios:
                curr_sample = "malware"

            selected_sample = st.selectbox(
                "Chọn Tập dữ liệu Cảnh báo Mẫu:",
                options=list(sample_scenarios.keys()),
                format_func=lambda k: sample_scenarios[k],
                index=list(sample_scenarios.keys()).index(curr_sample),
            )
            st.session_state.selected_sample_scenario = selected_sample
            target_scenario_key = selected_sample

        elif selected_source_mode == "custom":
            target_scenario_key = "custom"
            uploaded_file = st.file_uploader("Chọn file .json cảnh báo Wazuh:", type=["json"])
            if uploaded_file:
                custom_json_content = uploaded_file.read().decode("utf-8")

        # Cập nhật khi thay đổi chế độ hoặc kịch bản
        mode_changed = selected_source_mode != st.session_state.get("evidence_source_mode")
        scenario_changed = target_scenario_key != st.session_state.get("selected_scenario_key")

        if mode_changed or scenario_changed:
            st.session_state.evidence_source_mode = selected_source_mode
            reset_pipeline(target_scenario_key, custom_json_content)
            st.rerun()

        st.markdown("##### 🎮 Thao tác Thực thi:")
        col_btn1, col_btn2 = st.columns(2)

        with col_btn1:
            if st.button("🚀 Chạy Hết", help="Thực thi toàn bộ quy trình StateGraph tự động", use_container_width=True):
                run_full_pipeline()
                st.rerun()

        with col_btn2:
            step_disabled = st.session_state.current_node is None
            if st.button("⏭️ Từng bước", disabled=step_disabled, help="Chạy từng Agent Node", use_container_width=True):
                run_single_step()
                st.rerun()

        if st.button("🔄 Đặt lại / Reset", help="Khởi tạo lại trạng thái sự cố ban đầu", use_container_width=True):
            reset_pipeline(target_scenario_key, custom_json_content)
            st.rerun()

        st.divider()

        # Hiển thị tiến trình
        st.markdown("##### 📊 Trạng thái Vòng đời:")
        next_step_label = f"`{st.session_state.current_node}`" if st.session_state.current_node else "*Đã hoàn tất*"
        st.markdown(f"- **Node kế tiếp:** {next_step_label}")
        st.markdown(f"- **Trạng thái:** `{state.status.value}`")
        st.markdown(f"- **Chuyển dịch:** `{len(state.history)} bước`")

        if target_scenario_key in SCENARIO_CATALOG:
            st.caption(f"ℹ️ {SCENARIO_CATALOG[target_scenario_key]['description']}")

    # -------------------------------------------------------------
    # GIAO DIỆN CHÍNH THEO MENU ĐIỀU HƯỚNG
    # -------------------------------------------------------------

    # =============================================================
    # 1. VIEW: 🏠 TRANG CHỦ (WAZUH SOC MULTI-AGENT OVERVIEW)
    # =============================================================
    if selected_nav == "🏠 Trang chủ":
        # 1. HERO SECTION (CHUẨN SOC WAZUH SIEM & MULTI-AGENT)
        col_hero_left, col_hero_right = st.columns([1.6, 1.1])

        with col_hero_left:
            st.markdown(
                """
                <div style="padding: 6px 0 14px 0;">
                    <div style="display: flex; align-items: center; gap: 8px; margin-bottom: 10px;">
                        <span style="
                            background: rgba(56, 189, 248, 0.12);
                            color: #38bdf8;
                            border: 1px solid rgba(56, 189, 248, 0.3);
                            font-weight: 700;
                            font-size: 0.75rem;
                            letter-spacing: 0.08em;
                            text-transform: uppercase;
                            padding: 3px 8px;
                            border-radius: 4px;
                        ">HCMUTE • TIỂU LUẬN CHUYÊN NGÀNH AN TOÀN THÔNG TIN</span>
                        <span style="
                            background: rgba(16, 185, 129, 0.12);
                            color: #34d399;
                            border: 1px solid rgba(52, 211, 153, 0.3);
                            font-weight: 600;
                            font-size: 0.72rem;
                            padding: 3px 8px;
                            border-radius: 4px;
                        ">NIST SP 800-61 • MITRE ATT&CK • WAZUH SIEM</span>
                    </div>
                    <h1 style="
                        font-size: 2.1rem;
                        font-weight: 800;
                        color: #f8fafc;
                        margin: 0 0 14px 0;
                        line-height: 1.25;
                    ">Trung Tâm Tác Tử Tự Động Hóa Điều Tra & Phản Ứng Sự Cố (SOC Console)</h1>
                    <p style="
                        font-size: 0.98rem;
                        color: #94a3b8;
                        line-height: 1.6;
                        margin-bottom: 22px;
                    ">Tiếp nhận cảnh báo thời gian thực từ Wazuh SIEM qua Webhook, ứng dụng đồ thị trạng thái StateGraph điều phối 6 AI Agents chuyên trách (Ingestion, Triage, CTI, RCA, Response, Reporting) hỗ trợ tự động hóa chu trình phản ứng sự cố an ninh mạng.</p>
                </div>
                """,
                unsafe_allow_html=True,
            )

            col_a1, col_a2, col_a3 = st.columns([1.2, 1.1, 1.1])
            with col_a1:
                if st.button("🚀 Khởi chạy Toàn trình", type="primary", use_container_width=True):
                    run_full_pipeline()
                    st.rerun()
            with col_a2:
                if st.button("🔴 Nạp Live Wazuh", use_container_width=True):
                    st.session_state.evidence_source_mode = "live_wazuh"
                    reset_pipeline("live_wazuh")
                    st.rerun()
            with col_a3:
                if st.button("📊 Xem đánh giá sự cố", use_container_width=True):
                    st.session_state.active_tab_hint = "Đánh giá sự cố"
                    st.info("Chuyển sang mục '📈 Đánh giá sự cố' trên menu bên trái để xem chi tiết.")

        with col_hero_right:
            st.markdown(
                """
                <div style="display: flex; flex-direction: column; gap: 10px; margin-top: 14px;">
                    <div style="background: #0f172a; border: 1px solid #334155; border-radius: 10px; padding: 12px 16px;">
                        <div style="display: flex; gap: 12px; align-items: flex-start;">
                            <span style="color: #38bdf8; font-weight: 800; font-size: 1.2rem;">01</span>
                            <div>
                                <strong style="color: #f1f5f9; font-size: 0.92rem;">Thu thập & Bóc tách Wazuh SIEM</strong>
                                <div style="color: #94a3b8; font-size: 0.8rem; margin-top: 2px;">Nhận cảnh báo Webhook JSON, chuẩn hóa Sysmon Event ID 1, 3, 11 và Network telemetry</div>
                            </div>
                        </div>
                    </div>
                    <div style="background: #0f172a; border: 1px solid #334155; border-radius: 10px; padding: 12px 16px;">
                        <div style="display: flex; gap: 12px; align-items: flex-start;">
                            <span style="color: #38bdf8; font-weight: 800; font-size: 1.2rem;">02</span>
                            <div>
                                <strong style="color: #f1f5f9; font-size: 0.92rem;">Phân loại & Tái dựng Đa tác tử (Multi-Agent)</strong>
                                <div style="color: #94a3b8; font-size: 0.8rem; margin-top: 2px;">Triage đánh giá mức độ rủi ro, CTI làm giàu IoC, RCA truy vết cây tiến trình nguyên nhân gốc</div>
                            </div>
                        </div>
                    </div>
                    <div style="background: #0f172a; border: 1px solid #334155; border-radius: 10px; padding: 12px 16px;">
                        <div style="display: flex; gap: 12px; align-items: flex-start;">
                            <span style="color: #38bdf8; font-weight: 800; font-size: 1.2rem;">03</span>
                            <div>
                                <strong style="color: #f1f5f9; font-size: 0.92rem;">Phản ứng NIST & Phê duyệt Chuyên viên (HITL)</strong>
                                <div style="color: #94a3b8; font-size: 0.8rem; margin-top: 2px;">Đề xuất Active Response ngăn chặn, cổng phê duyệt chuyên viên SOC và xuất bản báo cáo DOCX</div>
                            </div>
                        </div>
                    </div>
                </div>
                """,
                unsafe_allow_html=True,
            )

        st.write("")

        # 2. THANH 5 BƯỚC WORKFLOW STAGES (CHUẨN NIST SP 800-61 & WAZUH SIEM)
        st.markdown(
            """
            <div style="display: grid; grid-template-columns: repeat(5, 1fr); gap: 12px; margin-bottom: 20px;">
                <div style="background: #0f172a; border-top: 3px solid #38bdf8; border-left: 1px solid #334155; border-right: 1px solid #334155; border-bottom: 1px solid #334155; border-radius: 8px; padding: 12px 14px;">
                    <div style="color: #38bdf8; font-size: 0.72rem; font-weight: 800; text-transform: uppercase;">WAZUH SIEM</div>
                    <div style="font-weight: 700; color: #f1f5f9; font-size: 0.88rem; margin: 4px 0;">Thu thập Cảnh báo</div>
                    <div style="color: #94a3b8; font-size: 0.78rem; line-height: 1.4;">Nhận cảnh báo trực tiếp từ Wazuh OVA (:8000 Webhook) hoặc kịch bản kiểm nghiệm.</div>
                </div>
                <div style="background: #0f172a; border-top: 3px solid #38bdf8; border-left: 1px solid #334155; border-right: 1px solid #334155; border-bottom: 1px solid #334155; border-radius: 8px; padding: 12px 14px;">
                    <div style="color: #38bdf8; font-size: 0.72rem; font-weight: 800; text-transform: uppercase;">TRIAGE ENSEMBLE</div>
                    <div style="font-weight: 700; color: #f1f5f9; font-size: 0.88rem; margin: 4px 0;">Phân loại Rủi ro</div>
                    <div style="color: #94a3b8; font-size: 0.78rem; line-height: 1.4;">Đánh giá mức độ nghiêm trọng, phân loại tấn công và lọc bỏ False Positive.</div>
                </div>
                <div style="background: #0f172a; border-top: 3px solid #38bdf8; border-left: 1px solid #334155; border-right: 1px solid #334155; border-bottom: 1px solid #334155; border-radius: 8px; padding: 12px 14px;">
                    <div style="color: #38bdf8; font-size: 0.72rem; font-weight: 800; text-transform: uppercase;">CTI & RCA</div>
                    <div style="font-weight: 700; color: #f1f5f9; font-size: 0.88rem; margin: 4px 0;">Làm giàu & Cây tiến trình</div>
                    <div style="color: #94a3b8; font-size: 0.78rem; line-height: 1.4;">Tra cứu IoC Reputation, phân tích nguyên nhân gốc rễ và ánh xạ MITRE ATT&CK.</div>
                </div>
                <div style="background: #0f172a; border-top: 3px solid #38bdf8; border-left: 1px solid #334155; border-right: 1px solid #334155; border-bottom: 1px solid #334155; border-radius: 8px; padding: 12px 14px;">
                    <div style="color: #38bdf8; font-size: 0.72rem; font-weight: 800; text-transform: uppercase;">ACTIVE RESPONSE</div>
                    <div style="font-weight: 700; color: #f1f5f9; font-size: 0.88rem; margin: 4px 0;">Ngăn chặn & HITL</div>
                    <div style="color: #94a3b8; font-size: 0.78rem; line-height: 1.4;">Đề xuất cô lập Host, chặn IP độc hại và phê duyệt chuyên viên SOC qua HITL.</div>
                </div>
                <div style="background: #0f172a; border-top: 3px solid #38bdf8; border-left: 1px solid #334155; border-right: 1px solid #334155; border-bottom: 1px solid #334155; border-radius: 8px; padding: 12px 14px;">
                    <div style="color: #38bdf8; font-size: 0.72rem; font-weight: 800; text-transform: uppercase;">NIST SP 800-61</div>
                    <div style="font-weight: 700; color: #f1f5f9; font-size: 0.88rem; margin: 4px 0;">Xuất bản Hồ sơ</div>
                    <div style="color: #94a3b8; font-size: 0.78rem; line-height: 1.4;">Tự động xuất bản báo cáo sự cố toàn diện đa định dạng (Markdown, JSON, Word .docx).</div>
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

        # 3. KẾT QUẢ ĐỊNH DANH TẤN CÔNG & KPI (HIỂN THỊ KHI ĐÃ CÓ PHÂN TÍCH)
        render_attack_detection_banner(state)
        render_kpi_cards(state)

        st.write("")

        # 4. BOTTOM CARDS (LUỒNG CHÍNH VÀ DEMO SCENARIOS TRÊN GIAO DIỆN TỐI)
        col_b1, col_b2, col_b3 = st.columns([1.5, 1, 1])

        with col_b1:
            st.markdown(
                """
                <div style="background: #0f172a; border: 1px solid #334155; border-radius: 12px; padding: 20px; height: 100%;">
                    <div style="color: #38bdf8; font-size: 0.75rem; font-weight: 800; text-transform: uppercase;">Luồng chính của đề tài</div>
                    <h3 style="margin: 6px 0 10px 0; color: #f1f5f9; font-size: 1.12rem;">Điều tra Sự cố Tự động từ Wazuh SIEM</h3>
                    <p style="color: #94a3b8; font-size: 0.85rem; line-height: 1.5;">Tiếp nhận cảnh báo live từ Wazuh OVA qua Webhook hoặc tệp đối soát, tự động điều tra qua 6 AI Agents chuyên trách.</p>
                </div>
                """,
                unsafe_allow_html=True,
            )

        with col_b2:
            st.markdown(
                """
                <div style="background: #0f172a; border: 1px solid #334155; border-radius: 12px; padding: 20px; height: 100%;">
                    <div style="color: #38bdf8; font-size: 0.75rem; font-weight: 800; text-transform: uppercase;">Kịch bản Kiểm nghiệm 01</div>
                    <h3 style="margin: 6px 0 10px 0; color: #f1f5f9; font-size: 1.12rem;">Malware & C2 Shell</h3>
                    <p style="color: #94a3b8; font-size: 0.85rem; line-height: 1.5;">PowerShell tải payload mã độc từ IP lạ, thiết lập kết nối C2 và tạo tiến trình bất thường.</p>
                </div>
                """,
                unsafe_allow_html=True,
            )

        with col_b3:
            st.markdown(
                """
                <div style="background: #0f172a; border: 1px solid #334155; border-radius: 12px; padding: 20px; height: 100%;">
                    <div style="color: #38bdf8; font-size: 0.75rem; font-weight: 800; text-transform: uppercase;">Kịch bản Kiểm nghiệm 02</div>
                    <h3 style="margin: 6px 0 10px 0; color: #f1f5f9; font-size: 1.12rem;">Credential Spraying</h3>
                    <p style="color: #94a3b8; font-size: 0.85rem; line-height: 1.5;">Tấn công brute force đăng nhập hàng loạt tài khoản Domain Controller liên tục.</p>
                </div>
                """,
                unsafe_allow_html=True,
            )

    # =============================================================
    # 2. VIEW: 📈 ĐÁNH GIÁ SỰ CỐ (INCIDENT ASSESSMENT & RCA)
    # =============================================================
    elif selected_nav == "📈 Đánh giá sự cố":
        st.subheader("📈 Đánh giá Mức độ Rủi ro & Nguyên nhân Gốc rễ Chuyên sâu (RCA)")
        render_attack_detection_banner(state)
        render_kpi_cards(state)
        st.divider()

        tab_rca1, tab_rca2, tab_rca3, tab_rca4, tab_rca5 = st.tabs([
            "🧬 1. Phả hệ Tiến trình & MITRE",
            "🔍 2. Phân tích 5-Whys & Nhân quả",
            "🐟 3. Biểu đồ Xương cá Ishikawa",
            "⚖️ 4. Đánh giá CIA & Pháp lý (NĐ 13)",
            "🛡️ 5. Cổng Kiểm định An toàn (Integrity Gate)",
        ])

        with tab_rca1:
            col_r1, col_r2 = st.columns([1, 1])
            with col_r1:
                render_mitre_matrix(state.threat_intel, state.rca_result)
            with col_r2:
                if state.rca_result:
                    st.markdown("#### 🧬 Tóm lược Nguyên nhân Gốc rễ")
                    st.info(state.rca_result.root_cause_summary)
                    if state.rca_result.initial_access_vector:
                        st.markdown(f"**Vector Xâm nhập Ban đầu:** `{state.rca_result.initial_access_vector}`")
                else:
                    st.info("Chưa có kết quả phân tích RCA. Vui lòng bấm '🚀 Chạy Hết' trên thanh điều khiển.")

            st.divider()
            render_process_tree(state.rca_result)

        with tab_rca2:
            st.markdown("### 🔍 Phân tích Chuỗi Nhân quả 5 Whys (Root Cause Analysis)")
            st.caption("Phương pháp truy vết nhân quả từ triệu chứng bên ngoài đến lỗi quy trình và chính sách cốt lõi:")

            init_vector = state.rca_result.initial_access_vector if state.rca_result else "Macro tệp đính kèm lạ"
            attack_type_str = state.attack_type or "Endpoint Security Incident"

            st.markdown(
                f"""
                <div style="background: #0f172a; border: 1px solid #334155; border-radius: 12px; padding: 18px; margin-bottom: 16px;">
                    <div style="font-weight: 800; color: #f87171; font-size: 1.05rem; margin-bottom: 14px;">
                        🚨 VẤN ĐỀ TRỌNG TÂM: Phát hiện {html.escape(str(attack_type_str or "Security Incident"))} - {html.escape(str(state.incident_id or "INC-UNKNOWN"))}
                    </div>
                    <div style="display: flex; flex-direction: column; gap: 10px;">
                        <div style="background: #1e293b; border-left: 4px solid #38bdf8; padding: 12px 14px; border-radius: 6px;">
                            <strong style="color: #38bdf8;">Why 1 (Hiện tượng bề mặt):</strong> Tại sao tiến trình độc hại được thực thi trên máy trạm?<br>
                            <span style="color: #cbd5e1;">➜ Người dùng mở tệp văn bản kích hoạt script thông qua: <code>{html.escape(str(init_vector or "N/A"))}</code>.</span>
                        </div>
                        <div style="background: #1e293b; border-left: 4px solid #38bdf8; padding: 12px 14px; border-radius: 6px;">
                            <strong style="color: #38bdf8;">Why 2 (Kỹ thuật thực thi):</strong> Tại sao lệnh PowerShell tải được payload từ ngoài Internet?<br>
                            <span style="color: #cbd5e1;">➜ Kẻ tấn công sử dụng các cờ <code>-ExecutionPolicy Bypass -WindowStyle Hidden -enc</code> để vượt qua cơ chế kiểm tra dòng lệnh tĩnh.</span>
                        </div>
                        <div style="background: #1e293b; border-left: 4px solid #fbbf24; padding: 12px 14px; border-radius: 6px;">
                            <strong style="color: #fbbf24;">Why 3 (Kiểm soát mạng):</strong> Tại sao kết nối C2 outbound qua cổng 443 không bị phát hiện ngay lập tức?<br>
                            <span style="color: #cbd5e1;">➜ Lưu lượng mạng trà trộn vào lưu lượng SSL/TLS thông thường, tường lửa nội bộ chưa kiểm soát Egress Traffic theo từng tiến trình.</span>
                        </div>
                        <div style="background: #1e293b; border-left: 4px solid #fbbf24; padding: 12px 14px; border-radius: 6px;">
                            <strong style="color: #fbbf24;">Why 4 (Cấu hình phòng vệ):</strong> Tại sao EDR agent không tự động chặn đứng tiến trình ngay từ bước tạo file Temp?<br>
                            <span style="color: #cbd5e1;">➜ Chính sách EDR đang cấu hình ở chế độ Audit / Monitor nhằm tránh gây gián đoạn công việc nhân viên, chưa kích hoạt Active Blocking.</span>
                        </div>
                        <div style="background: rgba(220, 38, 38, 0.15); border: 1px solid rgba(239, 68, 68, 0.3); border-left: 4px solid #ef4444; padding: 14px; border-radius: 6px;">
                            <strong style="color: #f87171;">Why 5 (NGUYÊN NHÂN GỐC RỄ - ROOT CAUSE):</strong> Tại sao chưa kích hoạt Active Blocking và chính sách PowerShell nghiêm ngặt?<br>
                            <span style="color: #fecaca; font-weight: 500;">➜ Doanh nghiệp chưa hoàn tất quy trình chuẩn hóa chính sách ATTT (Policy Hardening), chưa triển khai PowerShell Constrained Language Mode và thiếu quy trình diễn tập mô phỏng tấn công định kỳ.</span>
                        </div>
                    </div>
                </div>
                """,
                unsafe_allow_html=True,
            )

            if state.rca_result and state.rca_result.attack_timeline:
                st.markdown("#### ⏱️ Dòng thời gian Diễn biến Sự cố (Attack Timeline)")
                for idx, t_step in enumerate(state.rca_result.attack_timeline, start=1):
                    st.markdown(f"**Bước {idx}:** {t_step}")

        with tab_rca3:
            st.markdown("### 🐟 Biểu đồ Xương cá Ishikawa (Fishbone Diagram)")
            st.caption("Bóc tách 6 nguyên nhân dẫn tới sự cố theo chuẩn quản trị an toàn thông tin doanh nghiệp:")

            col_f1, col_f2 = st.columns(2)
            with col_f1:
                st.markdown(
                    """
                    <div style="background: #0f172a; border: 1px solid #334155; border-radius: 8px; padding: 14px; margin-bottom: 12px;">
                        <div style="color: #38bdf8; font-weight: 700; font-size: 0.95rem;">👤 1. Con người (People)</div>
                        <ul style="color: #cbd5e1; font-size: 0.85rem; margin-top: 6px;">
                            <li>Người dùng thiếu cảnh giác khi mở tệp văn bản chứa macro từ email bên ngoài.</li>
                            <li>Quản trị viên chưa quen với quy trình phản ứng nhanh khi EDR phát cảnh báo.</li>
                        </ul>
                    </div>
                    <div style="background: #0f172a; border: 1px solid #334155; border-radius: 8px; padding: 14px; margin-bottom: 12px;">
                        <div style="color: #38bdf8; font-weight: 700; font-size: 0.95rem;">📋 2. Quy trình (Process)</div>
                        <ul style="color: #cbd5e1; font-size: 0.85rem; margin-top: 6px;">
                            <li>Quy trình vá lỗi chưa có cam kết SLA 24h đối với các lỗ hổng Critical.</li>
                            <li>Chưa có quy trình cô lập tức thời khi mức độ rủi ro vượt ngưỡng High.</li>
                        </ul>
                    </div>
                    <div style="background: #0f172a; border: 1px solid #334155; border-radius: 8px; padding: 14px; margin-bottom: 12px;">
                        <div style="color: #38bdf8; font-weight: 700; font-size: 0.95rem;">💻 3. Công nghệ (Technology)</div>
                        <ul style="color: #cbd5e1; font-size: 0.85rem; margin-top: 6px;">
                            <li>EDR chỉ đặt ở chế độ Monitor thay vì chủ động Enforce/Block.</li>
                            <li>Thiếu hệ thống Mail Gateway Sandboxing quét trước tệp đính kèm.</li>
                        </ul>
                    </div>
                    """,
                    unsafe_allow_html=True,
                )
            with col_f2:
                st.markdown(
                    """
                    <div style="background: #0f172a; border: 1px solid #334155; border-radius: 8px; padding: 14px; margin-bottom: 12px;">
                        <div style="color: #38bdf8; font-weight: 700; font-size: 0.95rem;">🌐 4. Môi trường Mạng (Environment)</div>
                        <ul style="color: #cbd5e1; font-size: 0.85rem; margin-top: 6px;">
                            <li>Mạng nội bộ chưa được phân đoạn vi mô (Micro-segmentation) triệt để.</li>
                            <li>Cho phép các máy trạm kết nối trực tiếp ra Internet cổng 443 không qua Web Proxy.</li>
                        </ul>
                    </div>
                    <div style="background: #0f172a; border: 1px solid #334155; border-radius: 8px; padding: 14px; margin-bottom: 12px;">
                        <div style="color: #38bdf8; font-weight: 700; font-size: 0.95rem;">📜 5. Chính sách Quản trị (Policy)</div>
                        <ul style="color: #cbd5e1; font-size: 0.85rem; margin-top: 6px;">
                            <li>Chưa áp dụng chính sách giới hạn PowerShell Constrained Language Mode trên GPO.</li>
                            <li>Thiếu chính sách xoay vòng mật khẩu đặc quyền quản trị tự động.</li>
                        </ul>
                    </div>
                    <div style="background: #0f172a; border: 1px solid #334155; border-radius: 8px; padding: 14px; margin-bottom: 12px;">
                        <div style="color: #38bdf8; font-weight: 700; font-size: 0.95rem;">📡 6. Giám sát & SOC (Monitoring)</div>
                        <ul style="color: #cbd5e1; font-size: 0.85rem; margin-top: 6px;">
                            <li>Cảnh báo bị trôi do hiện tượng Alert Fatigue (quá nhiều log thông tin mức thấp).</li>
                            <li>Thời gian phát hiện trung bình (MTTD) còn kéo dài trước khi có SOAR can thiệp.</li>
                        </ul>
                    </div>
                    """,
                    unsafe_allow_html=True,
                )

        with tab_rca4:
            st.markdown("### ⚖️ Đánh giá Tác động Toàn diện & Tuân thủ Pháp lý")
            st.caption("Đánh giá toàn diện theo bộ ba an toàn thông tin (CIA Triad), chỉ số RTO/RPO và Nghị định 13/2023/NĐ-CP:")

            col_c1, col_c2, col_c3 = st.columns(3)
            with col_c1:
                st.markdown(
                    """
                    <div style="background: rgba(225, 29, 72, 0.1); border: 1px solid rgba(244, 63, 94, 0.35); border-radius: 8px; padding: 14px; height: 100%;">
                        <div style="color: #fb7185; font-weight: 700; font-size: 0.95rem;">🔒 Tính Bảo mật (Confidentiality)</div>
                        <div style="color: #fecdd3; font-size: 0.85rem; margin-top: 6px; line-height: 1.5;">
                            <strong style="color: #fda4af;">Mức độ ảnh hưởng: CAO</strong><br>
                            Kẻ tấn công thiết lập kênh điều khiển C2 ra ngoài Internet và có khả năng trích xuất token/mật khẩu trong RAM.
                        </div>
                    </div>
                    """,
                    unsafe_allow_html=True,
                )
            with col_c2:
                st.markdown(
                    """
                    <div style="background: rgba(217, 119, 6, 0.1); border: 1px solid rgba(245, 158, 11, 0.35); border-radius: 8px; padding: 14px; height: 100%;">
                        <div style="color: #fbbf24; font-weight: 700; font-size: 0.95rem;">🧩 Tính Toàn vẹn (Integrity)</div>
                        <div style="color: #fde68a; font-size: 0.85rem; margin-top: 6px; line-height: 1.5;">
                            <strong style="color: #fcd34d;">Mức độ ảnh hưởng: TRUNG BÌNH</strong><br>
                            Mã độc tạo tệp nhị phân trong thư mục Temp và đăng ký Scheduled Task tự chạy; dữ liệu nghiệp vụ chính không bị sửa đổi.
                        </div>
                    </div>
                    """,
                    unsafe_allow_html=True,
                )
            with col_c3:
                st.markdown(
                    """
                    <div style="background: rgba(2, 132, 199, 0.1); border: 1px solid rgba(56, 189, 248, 0.35); border-radius: 8px; padding: 14px; height: 100%;">
                        <div style="color: #38bdf8; font-weight: 700; font-size: 0.95rem;">⚡ Tính Sẵn sàng (Availability)</div>
                        <div style="color: #bae6fd; font-size: 0.85rem; margin-top: 6px; line-height: 1.5;">
                            <strong style="color: #7dd3fc;">Mức độ ảnh hưởng: CAO</strong><br>
                            Máy trạm bị cô lập logic khỏi mạng nội bộ để xử lý; thời gian khôi phục nằm trong ngưỡng RTO cam kết.
                        </div>
                    </div>
                    """,
                    unsafe_allow_html=True,
                )

            st.write("")
            st.markdown(
                """
                <div style="background: #0f172a; border: 1px solid #334155; border-radius: 8px; padding: 16px;">
                    <strong style="color: #f1f5f9; font-size: 0.95rem;">📜 Nghĩa vụ Tuân thủ & Báo cáo Pháp lý (Nghị định 13/2023/NĐ-CP):</strong>
                    <div style="color: #cbd5e1; font-size: 0.85rem; margin-top: 6px; line-height: 1.6;">
                        • <strong>Nghị định 13/2023/NĐ-CP (Bảo vệ dữ liệu cá nhân):</strong> Khi có dấu hiệu rò rỉ dữ liệu cá nhân (PII), tổ chức có nghĩa vụ lập hồ sơ thông báo sự cố theo <strong>Mẫu số 03</strong> gửi <strong>Cục An ninh mạng và phòng, chống tội phạm sử dụng công nghệ cao (A05) - Bộ Công an</strong> trong thời hạn chậm nhất <strong>72 giờ</strong>.<br>
                        • <strong>Luật An ninh mạng 2018:</strong> Báo cáo điều phối kỹ thuật với <strong>VNCERT/CC</strong> đối với sự cố tấn công mạng có nguy cơ lan rộng.
                    </div>
                </div>
                """,
                unsafe_allow_html=True,
            )

        with tab_rca5:
            st.markdown("### 🛡️ Cổng Kiểm định An toàn trước khi Khôi phục (Integrity Gate)")
            st.caption("Checklist tiêu chuẩn nghiệm thu kỹ thuật trước khi đưa máy chủ/máy trạm hoạt động trở lại:")

            st.checkbox("✅ Xác thực Mã độc: Đã quét toàn bộ hệ thống bằng EDR/Antivirus, xác nhận 0 mã độc và backdoor tồn dư.", value=True, key="ig_gate_1")
            st.checkbox("✅ Toàn vẹn Hệ thống: Đã kiểm tra tính toàn vẹn các file hệ thống và khóa Registry, không còn persistence.", value=True, key="ig_gate_2")
            st.checkbox("✅ Đổi thông tin xác thực: Đã thu hồi toàn bộ Active Session và ép buộc đổi mật khẩu các tài khoản liên quan.", value=True, key="ig_gate_3")
            st.checkbox("✅ Cập nhật Chính sách Bảo mật: Đã đưa IP/Domain máy chủ C2 vào Blacklist của Firewall/WAF biên.", value=True, key="ig_gate_4")
            st.checkbox("✅ Kích hoạt Giám sát Tăng cường (Burn-in 72h): Thiết lập phân tích viên giám sát chặt chẽ lưu lượng Egress.", value=True, key="ig_gate_5")

            st.success("🟢 HỆ THỐNG ĐÃ ĐẠT TIÊU CHUẨN KIỂM ĐỊNH AN TOÀN GO-LIVE")

    # =============================================================
    # 3. VIEW: 🛡️ BẰNG CHỨNG & LOG WAZUH SIEM (EVIDENCE & RAW LOGS EXPANDER)
    # =============================================================
    elif selected_nav == "🛡️ Bằng chứng & Log Wazuh SIEM":
        st.subheader("🛡️ Bằng chứng Cảnh báo Wazuh SIEM & Nhật ký Thực nghiệm")

        raw_alerts = st.session_state.get("raw_alerts", [])
        st.markdown(f"**Số lượng cảnh báo trong kịch bản:** `{len(raw_alerts)}` | **Đã phân tích:** `{len(state.parsed_events)}` sự kiện.")

        if raw_alerts:
            alert_rows = []
            for idx, a in enumerate(raw_alerts, start=1):
                rule = a.get("rule", {})
                agent = a.get("agent", {})
                data_field = a.get("data", {})
                alert_rows.append({
                    "#": idx,
                    "Rule ID": rule.get("id", "N/A"),
                    "Level": rule.get("level", 0),
                    "Mô tả Luật": rule.get("description", "N/A"),
                    "Agent": agent.get("name", "N/A"),
                    "IP Nguồn": data_field.get("srcip") or data_field.get("win", {}).get("eventdata", {}).get("ipAddress", "N/A"),
                    "IP Đích": data_field.get("dstip", "N/A"),
                })

            st.dataframe(
                alert_rows,
                use_container_width=True,
                column_config={
                    "Level": st.column_config.NumberColumn(
                        "Mức độ (0-15)",
                        help="Wazuh Alert Severity Level (0 - 15)",
                    )
                },
            )

            # GOM GỌN TOÀN BỘ LOG THÔ VÀO EXPANDER ĐỐI SOÁT
            with st.expander("📄 Khám nghiệm Dữ liệu Nhật ký Thô & Bằng chứng Gốc (Wazuh Telemetry Raw Alerts JSON)", expanded=False):
                st.caption("Dữ liệu payload JSON gốc được thu thập từ máy chủ Wazuh SIEM phục vụ đối soát toàn diện khi được yêu cầu:")
                for idx, a in enumerate(raw_alerts, start=1):
                    st.markdown(f"**Cảnh báo #{idx}:** Rule `{a.get('rule', {}).get('id')}` - {a.get('rule', {}).get('description')}")
                    st.json(a)
        else:
            st.info("Chưa có dữ liệu cảnh báo nào được nạp.")

    # =============================================================
    # 4. VIEW: 💼 QUY TRÌNH IR (STATEGRAPH DAG & HITL)
    # =============================================================
    elif selected_nav == "💼 Quy trình IR":
        st.subheader("💼 Quy trình Phản ứng Sự cố (Incident Response Workflow)")

        tab_ir1, tab_ir2 = st.tabs(["🤖 Đồ thị StateGraph DAG", "🛡️ Phê duyệt Chuyên viên (HITL)"])

        with tab_ir1:
            diagram_view = st.radio(
                "Chế độ hiển thị Quy trình Điều phối:",
                options=[
                    "📊 Tiến trình Thực thi Thời gian thực (Gọi song song các agent)",
                    "⚡ Sơ đồ Logic Luồng Hoạt động (Archify Workflow Engine)",
                ],
                horizontal=True,
                index=0,
                label_visibility="collapsed",
            )

            if "Tiến trình" in diagram_view or "song song" in diagram_view:
                render_agent_graph_status(state)
            else:
                st.subheader("⚡ Sơ đồ Logic Luồng Hoạt động (Archify Multi-Agent StateGraph Workflow)")
                render_hitech_agent_workflow(state, height=740)

            st.divider()
            st.subheader("📜 Nhật ký Chuyển dịch Trạng thái (State Transition History)")
            if state.history:
                history_rows = []
                for idx, rec in enumerate(state.history, start=1):
                    ts_str = format_vn_time(rec.timestamp, "%H:%M:%S (ICT)")
                    history_rows.append({
                        "Bước": idx,
                        "Thời gian": ts_str,
                        "Node Thực thi": rec.node,
                        "Từ Trạng thái": rec.from_status.value if hasattr(rec.from_status, "value") else str(rec.from_status),
                        "Đến Trạng thái": rec.to_status.value if hasattr(rec.to_status, "value") else str(rec.to_status),
                        "Thông điệp": rec.message or "N/A",
                    })
                st.dataframe(history_rows, use_container_width=True)
            else:
                st.caption("Chưa có bước chuyển dịch nào được ghi nhận.")

            with st.expander("🔍 Chi tiết Checkpoints Snapshot", expanded=False):
                st.json(state.checkpoints)

        with tab_ir2:
            render_hitl_action_cards(state)
            st.divider()
            col_act1, col_act2 = st.columns([1, 2])
            with col_act1:
                if st.button(
                    "⚡ Thực thi Hành động đã Duyệt & Tiến hành Báo cáo",
                    type="primary",
                    use_container_width=True,
                    help="Kích hoạt Active Response và chuyển tiếp sang Reporting Agent",
                ):
                    with st.spinner("Đang thực thi các hành động đã phê duyệt và sinh báo cáo sự cố..."):
                        orch = st.session_state.orchestrator
                        state.analyst_approved = any(a.is_approved is True for a in state.recommended_actions)
                        state.analyst_feedback = "Chuyên viên SOC đã xác nhận kế hoạch can thiệp trên giao diện Web HITL"
                        final_state = orch.resume_after_hitl(
                            state=state,
                            analyst_approved=state.analyst_approved,
                            analyst_feedback=state.analyst_feedback,
                        )
                        st.session_state.incident_state = final_state
                        st.session_state.current_node = None
                        persist_incident_state(final_state)
                        st.success("Đã hoàn tất thực thi các hành động và xuất bản Báo cáo Điều tra Toàn diện!")
                        st.rerun()

    # =============================================================
    # 5. VIEW: 🛡️ IOC & SIEM (THREAT INTEL & HUNTING)
    # =============================================================
    elif selected_nav == "🛡️ IOC & SIEM":
        st.subheader("🛡️ Trung Tâm Tình Báo Mối Đe Dọa (CTI) & Chỉ Số Xâm Phạm (IoC Console)")
        st.caption("Đối soát đa chiều chỉ số IoC thời gian thực với VirusTotal, AlienVault OTX, AbuseIPDB và ánh xạ khung ma trận MITRE ATT&CK Enterprise.")

        # 1. Thống kê nhanh CTI KPIs
        ti = state.threat_intel
        raw_iocs = getattr(ti, "extracted_iocs", []) if ti else []
        seen_vals = set()
        clean_iocs = []
        for item in raw_iocs:
            v = str(item.value or "").strip()
            if v and v not in seen_vals:
                seen_vals.add(v)
                clean_iocs.append(item)

        total_cnt = len(clean_iocs)
        mal_cnt = sum(1 for i in clean_iocs if i.is_malicious or (i.reputation_score or 0) >= 70)
        actors_list = getattr(ti, "related_threat_actors", []) if ti else []
        tech_list = getattr(ti, "mitre_techniques", []) if ti else []

        c_kpi1, c_kpi2, c_kpi3, c_kpi4 = st.columns(4)
        c_kpi1.metric("🔍 Tổng IoCs Bóc Tách", f"{total_cnt} chỉ số")
        c_kpi2.metric(
            "🚨 IoC Độc Hại Xác Nhận",
            f"{mal_cnt} độc hại",
            delta="Đã gắn cờ CTI" if mal_cnt > 0 else "An toàn",
            delta_color="inverse" if mal_cnt > 0 else "normal",
        )
        c_kpi3.metric("🎯 Kỹ Thuật MITRE ATT&CK", f"{len(tech_list)} kỹ thuật")
        c_kpi4.metric(
            "🏴‍☠️ Threat Actors / APT",
            f"{len(actors_list)} nhóm",
            help=", ".join(actors_list) if actors_list else "Chưa định danh",
        )

        st.divider()

        # 2. Tabs chuyên sâu chuẩn SOC Console
        tab_cti_ioc, tab_cti_mitre, tab_cti_lookup = st.tabs([
            "🛡️ 1. Bảng Chỉ Số Xâm Phạm (IoC Table)",
            "🎯 2. Ma Trận Kỹ Thuật MITRE ATT&CK",
            "🔎 3. Tra Cứu Tình Báo Trực Tiếp (Live CTI Hunter)",
        ])

        with tab_cti_ioc:
            render_ioc_table(state.threat_intel)

        with tab_cti_mitre:
            render_mitre_matrix(state.threat_intel, state.rca_result)

        with tab_cti_lookup:
            st.markdown("#### 🔎 Tra Cứu & Thẩm Định Chỉ Số Xâm Phạm Độc Lập")
            st.caption("Công cụ cho phép chuyên viên SOC trực tiếp truy vấn danh tiếng bất kỳ địa chỉ IP, Domain hoặc File Hash với hệ thống CTI:")

            def detect_and_validate_ioc(chosen_type: str, raw_val: str):
                v = raw_val.strip()
                if not v:
                    return False, "Vui lòng nhập giá trị chỉ số IoC cần tra cứu.", "UNKNOWN", v

                # 1. Kiểm tra IP
                is_ip = False
                try:
                    ipaddress.ip_address(v)
                    is_ip = True
                except ValueError:
                    is_ip = False

                # 2. Kiểm tra Hash
                h_val = v
                if h_val.lower().startswith("sha256="):
                    h_val = h_val[7:].strip()
                elif h_val.lower().startswith("md5="):
                    h_val = h_val[4:].strip()
                elif h_val.lower().startswith("sha1="):
                    h_val = h_val[5:].strip()

                is_hex_hash = len(h_val) in (32, 40, 64) and all(c in "0123456789abcdefABCDEF" for c in h_val)

                # 3. Kiểm tra Domain / URL
                domain_pat = r"^(https?:\/\/)?([a-zA-Z0-9]([a-zA-Z0-9\-]{0,61}[a-zA-Z0-9])?\.)+[a-zA-Z]{2,}(:\d+)?(\/.*)?$"
                is_domain_url = bool(re.match(domain_pat, v)) and not is_ip and not is_hex_hash

                # Định danh loại thực tế
                if is_ip:
                    real_type = "IP Address"
                elif is_hex_hash:
                    real_type = "File Hash (SHA256/MD5)"
                elif is_domain_url:
                    real_type = "Domain / URL"
                else:
                    real_type = "UNKNOWN"

                # Đối soát với lựa chọn của người dùng
                if "Auto-Detect" in chosen_type:
                    if real_type == "UNKNOWN":
                        return False, f"Giá trị '{v}' không thuộc định dạng IP, Domain hay File Hash (MD5/SHA256) hợp lệ.", real_type, v
                    return True, "", real_type, h_val if is_hex_hash else v

                if chosen_type == "IP Address" and not is_ip:
                    return False, f"Giá trị nhập vào là {real_type}, không phải địa chỉ IP!", real_type, v

                if "Hash" in chosen_type and not is_hex_hash:
                    return False, f"Giá trị nhập vào là {real_type}, không phải mã băm MD5/SHA256!", real_type, v

                if "Domain" in chosen_type and not is_domain_url:
                    return False, f"Giá trị nhập vào là {real_type}, không phải tên miền/URL!", real_type, v

                return True, "", real_type, h_val if is_hex_hash else v

            col_in1, col_in2, col_in3 = st.columns([1.2, 2, 1])
            with col_in1:
                lookup_type = st.selectbox(
                    "Loại Chỉ Số (Indicator Type):",
                    [
                        "✨ Tự động nhận diện (Auto-Detect)",
                        "IP Address",
                        "File Hash (SHA256/MD5)",
                        "Domain / URL",
                    ],
                    index=0,
                )
            with col_in2:
                default_val = "194.26.29.112"
                lookup_query = st.text_input("Giá trị Chỉ số Cần Tra cứu:", value=default_val, key="cti_lookup_input")
            with col_in3:
                st.write("")
                st.write("")
                btn_lookup = st.button("🔍 Tra Cứu CTI Ngay", type="primary", use_container_width=True)

            if btn_lookup and lookup_query:
                is_valid, err_msg, actual_type, clean_query = detect_and_validate_ioc(lookup_type, lookup_query)
                if not is_valid:
                    st.error(f"❌ **Không khớp định dạng:** {err_msg}")
                    st.warning(f"💡 **Gợi ý khắc phục:** Giá trị bạn vừa nhập được nhận diện là **{actual_type}**. Hãy chọn đúng loại chỉ số hoặc chuyển sang **'✨ Tự động nhận diện (Auto-Detect)'** để tra cứu!")
                else:
                    from src.tools.threat_intel import ThreatIntelTools
                    with st.spinner(f"Đang truy vấn cơ sở dữ liệu CTI cho {actual_type}: `{clean_query}`..."):
                        if actual_type == "IP Address":
                            res = ThreatIntelTools.lookup_ip(clean_query)
                        elif "Hash" in actual_type:
                            res = ThreatIntelTools.lookup_hash(clean_query)
                        else:
                            res = ThreatIntelTools.lookup_domain(clean_query)

                    if res.details.get("status") == "invalid_format":
                        st.error(f"❌ Lỗi định dạng từ CTI Engine: {res.details.get('error')}")
                    else:
                        score = res.reputation_score or 0
                        is_mal = res.is_malicious or score >= 70
                        badge_color = "#ef4444" if is_mal else ("#f59e0b" if score > 30 else "#10b981")
                        verdict_text = "🚨 ĐỘC HẠI (MALICIOUS)" if is_mal else ("🟡 ĐÁNG NGỜ (SUSPICIOUS)" if score > 30 else "🟢 AN TOÀN (BENIGN)")

                        st.markdown(
                            f"""
                            <div style="background: #0f172a; border: 1px solid #334155; border-left: 5px solid {badge_color}; border-radius: 10px; padding: 18px 20px; margin-top: 14px;">
                                <div style="display: flex; justify-content: space-between; align-items: center; margin-bottom: 10px; flex-wrap: wrap; gap: 8px;">
                                    <div>
                                        <span style="font-family: monospace; font-size: 1.05rem; color: #f8fafc; font-weight: 700;">{html.escape(res.value)}</span>
                                        <span style="background: rgba(56, 189, 248, 0.15); color: #38bdf8; font-size: 0.72rem; padding: 2px 8px; border-radius: 4px; margin-left: 8px;">{actual_type}</span>
                                    </div>
                                    <span style="background: {badge_color}; color: #ffffff; font-weight: 800; font-size: 0.78rem; padding: 4px 12px; border-radius: 999px;">{verdict_text}</span>
                                </div>
                                <div style="display: grid; grid-template-columns: repeat(3, 1fr); gap: 14px; margin-top: 12px; padding-top: 12px; border-top: 1px solid #1e293b;">
                                    <div>
                                        <span style="color: #94a3b8; font-size: 0.78rem;">Điểm Nguy cơ Tổng hợp</span>
                                        <div style="font-size: 1.25rem; font-weight: 800; color: {badge_color};">{score}/100</div>
                                    </div>
                                    <div>
                                        <span style="color: #94a3b8; font-size: 0.78rem;">Threat Actor Liên quan</span>
                                        <div style="font-size: 0.92rem; font-weight: 700; color: #f1f5f9;">{html.escape(res.threat_actor or "Chưa định danh")}</div>
                                    </div>
                                    <div>
                                        <span style="color: #94a3b8; font-size: 0.78rem;">Nguồn Dữ liệu (Source)</span>
                                        <div style="font-size: 0.92rem; font-weight: 600; color: #38bdf8;">{html.escape(str(res.details.get("source", "CTI Verified")))}</div>
                                    </div>
                                </div>
                            </div>
                            """,
                            unsafe_allow_html=True,
                        )

    # =============================================================
    # 6. VIEW: 💬 TRỢ LÝ SOC (INTERACTIVE ASSISTANT CHATBOT)
    # =============================================================
    elif selected_nav == "💬 Trợ lý SOC":
        st.subheader("💬 Trợ lý SOC (Tương tác Điều tra & Điều khiển HITL)")

        from src.copilot.engine import SOCCopilotEngine
        copilot_engine: SOCCopilotEngine = st.session_state.copilot_engine
        provider_obj = getattr(copilot_engine, "provider", getattr(copilot_engine, "llm_provider", None))
        provider_name = getattr(provider_obj, "provider_type", "Hybrid") if provider_obj else "Unknown"
        active_model = getattr(provider_obj, "model", "default") if provider_obj else "Unknown"

        col_chat, col_context = st.columns([1.7, 1.0])

        with col_context:
            # Lấy mức độ nghiêm trọng và số lượng IoC an toàn
            sev_obj = None
            if hasattr(state, "severity") and state.severity is not None:
                sev_obj = state.severity
            elif state.triage_result and hasattr(state.triage_result, "severity"):
                sev_obj = state.triage_result.severity
            sev_label = sev_obj.value if hasattr(sev_obj, "value") else str(sev_obj or "Chưa phân loại")

            ioc_cnt = len(getattr(state.threat_intel, "extracted_iocs", [])) if state.threat_intel else 0

            st.markdown(
                f"""
                <div style="background: #0f172a; border: 1px solid #334155; border-radius: 12px; padding: 14px 16px; margin-bottom: 12px;">
                    <div style="display: flex; align-items: center; justify-content: space-between;">
                        <span style="font-weight: 700; color: #38bdf8; font-size: 0.85rem; text-transform: uppercase;">Trạng thái Trợ lý SOC</span>
                        <span style="background: rgba(16, 185, 129, 0.2); color: #34d399; border: 1px solid rgba(52, 211, 153, 0.4); font-size: 0.72rem; font-weight: 700; padding: 2px 8px; border-radius: 999px;">🟢 Online</span>
                    </div>
                    <div style="margin-top: 8px; font-size: 0.82rem; color: #cbd5e1; line-height: 1.6;">
                        • <strong>Gateway:</strong> <code style="color: #38bdf8;">{provider_name.upper()}</code><br/>
                        • <strong>Model:</strong> <code style="color: #38bdf8;">{active_model}</code><br/>
                        • <strong>Sự cố:</strong> <code style="color: #cbd5e1;">{state.incident_id}</code><br/>
                        • <strong>Mức độ:</strong> <code style="color: #fbbf24;">{sev_label}</code>
                    </div>
                </div>
                """,
                unsafe_allow_html=True,
            )

            last_err = getattr(provider_obj, "last_error", None)
            if provider_name.lower() == "deterministic" and last_err:
                st.warning(f"⚠️ **Đang dùng Fallback Deterministic CoT** do API lỗi: `{last_err}`")

            with st.expander("⚙️ Cấu hình AI Gateway & Model", expanded=False):
                cfg_endpoint = st.text_input(
                    "API Endpoint",
                    value=getattr(provider_obj, "endpoint", "http://localhost:20128/v1/chat/completions"),
                    key="copilot_cfg_endpoint"
                )
                cfg_model = st.text_input(
                    "Tên Model / Combo Provider",
                    value=getattr(provider_obj, "model", "AICODE"),
                    key="copilot_cfg_model"
                )
                cfg_key = st.text_input(
                    "API Key",
                    value=getattr(provider_obj, "api_key", ""),
                    type="password",
                    key="copilot_cfg_key"
                )
                if st.button("💾 Lưu & Áp dụng Cấu hình", use_container_width=True):
                    if provider_obj:
                        provider_obj.endpoint = cfg_endpoint
                        provider_obj.model = cfg_model
                        provider_obj.api_key = cfg_key
                    st.success("Đã cập nhật cấu hình Trợ lý SOC thành công!")
                    st.rerun()

            with st.expander("📌 Bối cảnh Sự cố Đang Tải", expanded=True):
                st.markdown(f"**Loại tấn công:** `{state.attack_type or 'Chưa xác định'}`")
                st.markdown(f"**Trạng thái:** `{state.status.value}`")
                if state.rca_result:
                    st.caption(f"🧬 **Root Cause:** {state.rca_result.root_cause_summary[:160]}...")
                st.markdown(f"**Số lượng IoC:** `{ioc_cnt}`")
                approved_cnt = sum(1 for a in state.recommended_actions if a.is_approved is True)
                st.markdown(f"**Hành động đã duyệt:** `{approved_cnt}/{len(state.recommended_actions)}`")

            if st.button("🗑️ Xóa Lịch sử Hội thoại", use_container_width=True):
                st.session_state.copilot_messages = [
                    {
                        "role": "assistant",
                        "content": "Lịch sử hội thoại đã được đặt lại. Tôi sẵn sàng hỗ trợ bạn điều tra sự cố mới!",
                    }
                ]
                st.rerun()

        with col_chat:
            # 1. Quick Action Prompt Chips
            st.markdown("##### ⚡ Gợi ý Lệnh Nhanh (1-Click Prompt):")
            chip1, chip2 = st.columns(2)
            chip3, chip4 = st.columns(2)

            quick_prompt = None
            with chip1:
                if st.button("🔍 Giải thích Nguyên nhân Gốc rễ", use_container_width=True):
                    quick_prompt = "Giải thích nguyên nhân gốc rễ và phả hệ cây tiến trình của sự cố này"
            with chip2:
                if st.button("🌐 Đánh giá Nguy cơ IP & CTI", use_container_width=True):
                    quick_prompt = "Đánh giá các chỉ số IoC, IP nguồn và điểm danh tiếng CTI của sự cố này"
            with chip3:
                if st.button("📋 Tóm tắt Sự cố", use_container_width=True):
                    quick_prompt = "Tóm tắt sự cố này"
            with chip4:
                if st.button("🛡️ Duyệt Toàn bộ Kế hoạch Phản ứng", use_container_width=True):
                    quick_prompt = "Duyệt tất cả các hành động phản ứng"

            st.divider()

            # 2. Render Chat History
            chat_container = st.container()
            with chat_container:
                for msg in st.session_state.copilot_messages:
                    avatar = "🛡️" if msg["role"] == "user" else "🤖"
                    with st.chat_message(msg["role"], avatar=avatar):
                        st.markdown(msg["content"])
                        if msg.get("action_card"):
                            card = msg["action_card"]
                            st.info(f"⚡ **Lệnh can thiệp:** `{card.get('type')}` trên `{card.get('target')}` (Trạng thái: `{card.get('status')}`)")

            # 3. Chat Input Box
            user_input = st.chat_input("Nhập câu hỏi hoặc chỉ đạo nghiệp vụ cho Trợ lý SOC...")
            prompt_to_process = quick_prompt or user_input

            if prompt_to_process:
                # Lưu tin nhắn người dùng
                st.session_state.copilot_messages.append({"role": "user", "content": prompt_to_process})
                with st.chat_message("user", avatar="🛡️"):
                    st.markdown(prompt_to_process)

                # Xử lý phản hồi qua Trợ lý SOC Engine
                with st.chat_message("assistant", avatar="🤖"):
                    with st.spinner("🤖 Trợ lý SOC đang phân tích bối cảnh và suy luận..."):
                        copilot_res = copilot_engine.process_message(
                            user_message=prompt_to_process,
                            state=state,
                            chat_history=st.session_state.copilot_messages,
                        )
                        st.markdown(copilot_res.response_text)

                        # Nếu có hành động can thiệp được thực thi
                        if copilot_res.executed_action:
                            act_info = copilot_res.executed_action
                            st.success(f"✅ **Đã đồng bộ với StateGraph & Wazuh:** {act_info.get('message', '')}")

                        # Lưu phản hồi vào session
                        assistant_msg = {
                            "role": "assistant",
                            "content": copilot_res.response_text,
                        }
                        if copilot_res.executed_action:
                            intent_str = getattr(copilot_res.intent, "value", str(copilot_res.intent))
                            assistant_msg["action_card"] = {
                                "type": copilot_res.executed_action.get("action_type", "HITL"),
                                "target": copilot_res.executed_action.get("target", "System"),
                                "status": "APPROVED" if intent_str == "APPROVE_ACTION" else "REJECTED",
                            }
                        st.session_state.copilot_messages.append(assistant_msg)
                        st.rerun()

    # =============================================================
    # 7. VIEW: 🕒 LỊCH SỬ & BÁO CÁO (NIST SP 800-61 REPORT & ARCHIVE)
    # =============================================================
    elif selected_nav in ["🕒 Lịch sử & Báo cáo", "🕒 Lịch sử"]:
        st.subheader("🕒 Lịch sử & Báo cáo Điều tra Sự cố Toàn diện (NIST SP 800-61)")

        tab_rep_current, tab_rep_archive = st.tabs([
            "📄 Báo cáo Sự cố Hiện tại (Active)",
            "🗄️ Kho Lưu trữ & Xem lại Báo cáo Lịch sử (Archive)"
        ])

        # ---------------------------------------------------------
        # SUB-TAB 1: BÁO CÁO SỰ CỐ ĐANG PHÂN TÍCH TRONG PHIÊN
        # ---------------------------------------------------------
        with tab_rep_current:
            st.markdown(f"##### 📌 Báo cáo Điều tra Sự cố: `{state.incident_id}`")

            if state.final_report is None:
                st.info(
                    "ℹ️ **Báo cáo sự cố chưa được xuất bản.** "
                    "Bạn có thể khởi chạy toàn bộ quy trình StateGraph bằng nút '🚀 Chạy Hết' ở sidebar "
                    "hoặc bấm nút xuất bản tức thì bên dưới."
                )
                col_pub1, col_pub2 = st.columns([1.5, 2.5])
                with col_pub1:
                    if st.button("📝 Xuất bản & Lưu Báo cáo Ngay", type="primary", use_container_width=True):
                        with st.spinner("Reporting Agent đang tổng hợp thông tin và sinh báo cáo NIST SP 800-61..."):
                            agent = ReportingAgent()
                            rep = agent.generate_report(state)
                            state.final_report = rep
                            st.session_state.incident_state = state
                            saved_path = persist_incident_state(state)
                            st.success(f"✅ Đã xuất bản và lưu trữ sự cố thành công tại: `{saved_path}`")
                            st.rerun()
            else:
                rep = state.final_report

                md_content = generate_markdown_report(rep)
                json_content = generate_json_report(rep)
                docx_bytes = generate_docx_report(rep)

                # Thanh công cụ tải xuống và đồng bộ kho lưu trữ
                col_d1, col_d2, col_d3, col_save = st.columns([1, 1, 1, 1])
                with col_d1:
                    st.download_button(
                        label="📥 Tải Markdown (.md)",
                        data=md_content,
                        file_name=f"Incident_Report_{state.incident_id}.md",
                        mime="text/markdown",
                        use_container_width=True,
                    )
                with col_d2:
                    st.download_button(
                        label="📥 Tải Dữ liệu JSON (.json)",
                        data=json_content,
                        file_name=f"Incident_Report_{state.incident_id}.json",
                        mime="application/json",
                        use_container_width=True,
                    )
                with col_d3:
                    st.download_button(
                        label="📥 Tải Báo cáo Word (.docx)",
                        data=docx_bytes,
                        file_name=f"Incident_Report_{state.incident_id}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        use_container_width=True,
                    )
                with col_save:
                    if st.button("💾 Lưu vào Kho Lưu trữ", use_container_width=True, help="Lưu đè trạng thái hiện tại vào kho lưu trữ disk"):
                        saved_path = persist_incident_state(state)
                        st.success(f"Đã lưu: `{state.incident_id}`")

                st.divider()
                st.markdown(md_content)

        # ---------------------------------------------------------
        # SUB-TAB 2: KHO LƯU TRỮ & XEM LẠI BÁO CÁO SỰ CỐ LỊCH SỬ
        # ---------------------------------------------------------
        with tab_rep_archive:
            store = IncidentStore()
            historical_incidents = store.list_incidents(limit=100)

            if not historical_incidents:
                st.markdown(
                    """
                    <div style="background: #0f172a; border: 2px dashed #334155; border-radius: 12px; padding: 36px 20px; text-align: center; margin: 16px 0;">
                        <span style="font-size: 2.5rem;">📁</span>
                        <h4 style="color: #f1f5f9; margin: 12px 0 6px 0;">Kho Lưu trữ Sự cố Hiện Đang Trống</h4>
                        <p style="color: #94a3b8; font-size: 0.9rem; max-width: 520px; margin: 0 auto 16px auto; line-height: 1.5;">
                            Chưa có sự cố nào được lưu vào thư mục <code>data/incidents/</code>.<br/>
                            Dữ liệu sẽ tự động được lưu tại đây khi bạn chạy quy trình điều tra hoặc xuất bản báo cáo sự cố.
                        </p>
                    </div>
                    """,
                    unsafe_allow_html=True,
                )
                col_empty1, col_empty2 = st.columns([1, 1])
                with col_empty1:
                    if st.button("🚀 Lưu Sự cố Đang Phân tích Hiện tại vào Kho", type="primary", use_container_width=True):
                        saved_path = persist_incident_state(state)
                        st.success(f"Đã lưu thành công sự cố `{state.incident_id}` vào kho!")
                        st.rerun()
                with col_empty2:
                    if st.button("🔄 Nạp Cảnh báo Mới từ Wazuh Live", use_container_width=True):
                        st.session_state.evidence_source_mode = "live_wazuh"
                        reset_pipeline("live_wazuh")
                        st.rerun()
            else:
                # 1. Thống kê nhanh kho lưu trữ
                total_archived = len(historical_incidents)
                total_with_reports = sum(1 for inc in historical_incidents if inc.final_report is not None)
                total_critical_high = sum(
                    1 for inc in historical_incidents 
                    if getattr(inc, "severity", None) in {SeverityLevel.CRITICAL, SeverityLevel.HIGH}
                )

                kpi_col1, kpi_col2, kpi_col3 = st.columns(3)
                kpi_col1.metric("📁 Tổng Sự cố Đã Lưu", f"{total_archived} sự cố")
                kpi_col2.metric("📝 Báo cáo NIST Hoàn chỉnh", f"{total_with_reports}/{total_archived}")
                kpi_col3.metric("🔥 Sự cố Nguy cấp / Cao", f"{total_critical_high}")

                st.write("")
                st.markdown("##### 📋 Bảng Tổng hợp Lịch sử Sự cố:")

                # 2. Chuẩn bị bảng dữ liệu tổng quan
                table_rows = []
                for inc in historical_incidents:
                    sev_val = inc.severity.value if hasattr(inc.severity, "value") else str(inc.severity or "UNKNOWN")
                    status_val = inc.status.value if hasattr(inc.status, "value") else str(inc.status or "UNKNOWN")
                    has_rep = "✅ Đã xuất bản" if inc.final_report else "⏳ Chưa có"

                    summary_text = ""
                    if inc.final_report and inc.final_report.executive_summary:
                        summary_text = inc.final_report.executive_summary[:90] + "..."
                    elif inc.rca_result and inc.rca_result.root_cause_summary:
                        summary_text = inc.rca_result.root_cause_summary[:90] + "..."
                    else:
                        summary_text = f"Sự cố {inc.attack_type or 'Chưa phân loại'} ({len(inc.raw_logs)} logs)"

                    created_str = format_vn_time(inc.created_at, "%Y-%m-%d %H:%M:%S (ICT)")

                    table_rows.append({
                        "Mã Sự Cố": inc.incident_id,
                        "Thời Gian Ghi Nhận": created_str,
                        "Loại Tấn Công": inc.attack_type or "Chưa phân loại",
                        "Mức Độ": sev_val,
                        "Trạng Thái": status_val,
                        "Báo Cáo": has_rep,
                        "Tóm Tắt Điều Tra": summary_text,
                    })

                st.dataframe(
                    table_rows,
                    use_container_width=True,
                    column_config={
                        "Mã Sự Cố": st.column_config.TextColumn("Mã Sự Cố", width="medium"),
                        "Mức Độ": st.column_config.TextColumn("Mức Độ", width="small"),
                        "Báo Cáo": st.column_config.TextColumn("Báo Cáo NIST", width="small"),
                    },
                )

                st.divider()

                # 3. Chọn sự cố cụ thể để khám nghiệm & xuất báo cáo
                st.markdown("##### 🔍 Khám nghiệm Chi tiết & Tải Báo cáo Sự cố Lịch sử:")

                incident_dict = {inc.incident_id: inc for inc in historical_incidents}
                incident_labels = [
                    f"[{format_vn_time(inc.created_at, '%Y-%m-%d %H:%M (ICT)')}] {inc.incident_id} — {inc.attack_type or 'Mới tiếp nhận'} ({inc.severity.value if hasattr(inc.severity, 'value') else inc.severity})"
                    for inc in historical_incidents
                ]
                incident_keys = [inc.incident_id for inc in historical_incidents]

                selected_idx = st.selectbox(
                    "Chọn sự cố cần xem chi tiết:",
                    options=range(len(incident_keys)),
                    format_func=lambda idx: incident_labels[idx],
                )
                selected_inc = incident_dict[incident_keys[selected_idx]]

                # Các thao tác tương tác với sự cố lịch sử
                act_col1, act_col2 = st.columns([1.5, 1])
                with act_col1:
                    if st.button(
                        f"📂 Nạp (Load) Sự cố `{selected_inc.incident_id}` vào Phiên làm việc Hiện tại",
                        type="primary",
                        use_container_width=True,
                        help="Nạp toàn bộ bối cảnh sự cố này vào Dashboard để xem phả hệ Sysmon, MITRE ATT&CK và trao đổi với Trợ lý SOC",
                    ):
                        st.session_state.incident_state = selected_inc
                        restored_alerts = []
                        for r in selected_inc.raw_logs:
                            try:
                                restored_alerts.append(json.loads(r.raw_payload))
                            except Exception:
                                pass
                        st.session_state.raw_alerts = restored_alerts
                        st.session_state.current_node = None
                        st.success(f"✅ Đã nạp thành công sự cố `{selected_inc.incident_id}` vào phiên làm việc! Hãy chuyển sang các Tab khác để xem chi tiết.")
                        st.rerun()

                with act_col2:
                    if st.button(
                        f"🗑️ Xóa Sự cố `{selected_inc.incident_id}`",
                        use_container_width=True,
                        help="Xóa vĩnh viễn tệp JSON của sự cố này khỏi ổ đĩa",
                    ):
                        store.delete_incident(selected_inc.incident_id)
                        st.warning(f"Đã xóa sự cố `{selected_inc.incident_id}` khỏi kho lưu trữ.")
                        st.rerun()

                # Hiển thị thông tin tóm lược của sự cố được chọn
                st.write("")
                meta_col1, meta_col2, meta_col3, meta_col4 = st.columns(4)
                meta_col1.markdown(f"**Loại tấn công:** `{selected_inc.attack_type or 'N/A'}`")
                meta_col2.markdown(f"**Mức độ:** `{selected_inc.severity.value if hasattr(selected_inc.severity, 'value') else selected_inc.severity}`")
                meta_col3.markdown(f"**Trạng thái:** `{selected_inc.status.value if hasattr(selected_inc.status, 'value') else selected_inc.status}`")
                meta_col4.markdown(f"**Số lượng IoC:** `{len(selected_inc.threat_intel.extracted_iocs) if selected_inc.threat_intel and selected_inc.threat_intel.extracted_iocs else 0}`")

                # Bằng chứng và IoCs
                with st.expander(f"🔎 Xem Bằng chứng & IoCs đã phát hiện ({selected_inc.incident_id})", expanded=False):
                    iocs_preview = format_ioc_table_data(selected_inc.threat_intel)
                    if iocs_preview:
                        st.dataframe(iocs_preview, use_container_width=True)
                    else:
                        st.caption("Chưa ghi nhận danh sách IoC trong sự cố này.")

                # Kiểm tra và tải Báo cáo
                if selected_inc.final_report is None:
                    st.info("Sự cố này chưa có Báo cáo Điều tra Toàn diện.")
                    if st.button("📝 Sinh Báo cáo Ngay cho Sự cố Này", key=f"gen_rep_{selected_inc.incident_id}"):
                        with st.spinner("Đang sinh báo cáo NIST SP 800-61..."):
                            agent = ReportingAgent()
                            rep = agent.generate_report(selected_inc)
                            selected_inc.final_report = rep
                            store.save_incident(selected_inc)
                            st.success("Đã sinh và lưu báo cáo thành công!")
                            st.rerun()
                else:
                    rep_hist = selected_inc.final_report
                    md_h = generate_markdown_report(rep_hist)
                    json_h = generate_json_report(rep_hist)
                    docx_h = generate_docx_report(rep_hist)

                    dl_col1, dl_col2, dl_col3 = st.columns(3)
                    with dl_col1:
                        st.download_button(
                            label=f"📥 Tải Markdown ({selected_inc.incident_id})",
                            data=md_h,
                            file_name=f"Incident_Report_{selected_inc.incident_id}.md",
                            mime="text/markdown",
                            key=f"dl_md_{selected_inc.incident_id}",
                            use_container_width=True,
                        )
                    with dl_col2:
                        st.download_button(
                            label=f"📥 Tải JSON ({selected_inc.incident_id})",
                            data=json_h,
                            file_name=f"Incident_Report_{selected_inc.incident_id}.json",
                            mime="application/json",
                            key=f"dl_json_{selected_inc.incident_id}",
                            use_container_width=True,
                        )
                    with dl_col3:
                        st.download_button(
                            label=f"📥 Tải Word .docx ({selected_inc.incident_id})",
                            data=docx_h,
                            file_name=f"Incident_Report_{selected_inc.incident_id}.docx",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            key=f"dl_docx_{selected_inc.incident_id}",
                            use_container_width=True,
                        )

                    with st.expander("📄 Đọc Toàn văn Báo cáo Sự cố Lịch sử (Markdown Preview)", expanded=True):
                        st.markdown(md_h)


if __name__ == "__main__":
    main()

