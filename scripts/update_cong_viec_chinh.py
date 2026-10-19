# -*- coding: utf-8 -*-
"""
scripts/update_cong_viec_chinh.py
Cập nhật bảng tiến độ 13 tuần trong CongViecChinh.docx với các nội dung tích hợp Wazuh SIEM & EVE-NG.
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="D0D7DE", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def update_cong_viec_chinh():
    doc = docx.Document()
    
    # Thiết lập lề trang A4 Landscape để bảng rộng rãi, dễ đọc
    for section in doc.sections:
        section.page_width = Inches(11.69) # Landscape
        section.page_height = Inches(8.27)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    COLOR_PRIMARY = RGBColor(27, 54, 93)     # Deep Navy #1B365D
    COLOR_DARK = RGBColor(34, 34, 34)        # Charcoal #222222

    # Tiêu đề tài liệu
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(4)
    r1 = title_p.add_run("KẾ HOẠCH VÀ BẢNG PHÂN CÔNG NHIỆM VỤ CHI TIẾT (13 TUẦN)\n")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(16)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(14)
    r2 = sub_p.add_run("ĐỀ TÀI: XÂY DỰNG HỆ THỐNG MULTI-AGENT ỨNG DỤNG LLM HỖ TRỢ ĐIỀU TRA SỰ CỐ AN NINH MẠNG\n(Tích hợp Wazuh SIEM & Môi trường Giả lập EVE-NG)")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(12)
    r2.font.bold = True

    # Dữ liệu 13 tuần
    headers = ["Tuần", "Thời gian", "Hạng mục công việc trọng tâm", "Phân công chi tiết", "Sản phẩm bàn giao (Deliverable)"]
    widths = [Inches(0.6), Inches(1.2), Inches(2.6), Inches(3.3), Inches(2.4)]
    
    table_data = [
        ("1", "20/09 - 26/09", "Tiếp thu góp ý GVHD, chốt đề cương chi tiết", 
         "Cả nhóm (Nam: chuẩn hóa format tài liệu; Đô: rà soát mục tiêu kỹ thuật)", 
         "Đề cương chi tiết hoàn chỉnh"),
        
        ("2", "27/09 - 03/10", "Nghiên cứu nền tảng & Tổng quan", 
         "• Nam: Quy trình NIST SP 800-61, format log (Sysmon, Zeek, EVTX)\n• Đô: LLM, Prompt Engineering, Structured Output, Tool Calling", 
         "Bản thảo Báo cáo Chương 1 (Phần 1)"),
        
        ("3", "04/10 - 10/10", "Nghiên cứu Agent & Tri thức an ninh", 
         "• Nam: Threat Intelligence, MITRE ATT&CK (STIX/TAXII)\n• Đô: Kiến trúc LangGraph, State Management, Agent Orchestration", 
         "Bản thảo Báo cáo Chương 1 (Phần 2 & Khung Chương 2)"),
        
        ("4", "11/10 - 17/10", "Đặc tả yêu cầu & Kiến trúc hệ thống", 
         "• Cả nhóm: Xác định 3 kịch bản sự cố & Luồng dữ liệu chung\n• Đô: Thiết kế kiến trúc StateGraph điều phối\n• Nam: Thiết kế Schema Input/Output chuẩn (Pydantic v2)", 
         "Tài liệu System Architecture & Sequence Diagram"),
        
        ("5", "18/10 - 24/10", "Thiết kế chi tiết từng Agent & Tích hợp Wazuh (Làm song song)", 
         "• Nam: Thiết kế Ingestion Agent (kết nối Wazuh), Triage Agent (Wazuh Rule Level), Threat Intel Agent\n• Đô: Thiết kế RCA Agent, Response Agent (Wazuh Active Response), Reporting Agent", 
         "Bản đặc tả chi tiết Prompts, Schemas (WazuhAlert), Tools"),
        
        ("6", "25/10 - 31/10", "Dựng Lab EVE-NG, Hạ tầng Wazuh & Cơ sở tri thức RAG", 
         "• Nam: Thiết lập Topology EVE-NG (Firewall, Kali, Win10 Sysmon, Wazuh Server), xuất dataset alert thực tế\n• Đô: Xây dựng Vector DB/ChromaDB lưu trữ MITRE ATT&CK Matrix & Wazuh Ruleset", 
         "Hạ tầng Lab EVE-NG + Bộ dataset Wazuh alert thực tế + RAG Knowledge base sẵn sàng"),
        
        ("7", "01/11 - 07/11", "Hiện thực hóa Framework & Cụm Agent đầu vào", 
         "• Đô: Dựng khung LangGraph StateGraph, luồng chuyển trạng thái\n• Nam: Code Wazuh Ingestion Tool (API & File stream) & Detection/Triage Agent", 
         "Khung điều phối + Ingestion Agent tiếp nhận cảnh báo Wazuh chạy độc lập"),
        
        ("8", "08/11 - 14/11", "Hiện thực hóa Cụm Agent Chuyên sâu", 
         "• Nam: Code Threat Intel Agent (tích hợp API VirusTotal, AbuseIPDB)\n• Đô: Code RCA Agent (bóc tách tiến trình Sysmon từ Wazuh Alert, suy luận attack path)", 
         "2 Agent chuyên sâu vượt qua Unit Test"),
        
        ("9", "15/11 - 21/11", "Hiện thực Cụm Đề xuất/Báo cáo & TÍCH HỢP ĐỢT 1", 
         "• Đô: Code Response Agent (tích hợp Wazuh Active Response API) & Reporting Agent\n• Nam: Ráp cụm Ingestion-Triage-ThreatIntel vào LangGraph", 
         "Prototype v0.1: Chạy luồng từ Wazuh Alert → Phân tích đe dọa & Đề xuất phản ứng"),
        
        ("10", "22/11 - 28/11", "Tích hợp hoàn chỉnh & Dashboard UI", 
         "• Đô: Tích hợp RCA + Response + Report vào LangGraph, cơ chế HITL duyệt hành động\n• Nam: Xây dựng Dashboard UI bằng Streamlit hiển thị đồng thời Wazuh alerts & Agent graph", 
         "Prototype v1.0: Hoàn chỉnh toàn bộ luồng End-to-End có UI tương tác"),
        
        ("11", "29/11 - 05/12", "Thử nghiệm hệ thống trên 3 kịch bản EVE-NG", 
         "• Nam: Thực thi kịch bản Phishing & Scanning từ Kali EVE-NG, đo False Positive/Negative\n• Đô: Thực thi kịch bản Malware trên Windows 10 EVE-NG, đo thời gian điều tra (MTTI)", 
         "Bộ số liệu thực nghiệm đầy đủ trên môi trường mạng EVE-NG"),
        
        ("12", "06/12 - 12/12", "Đánh giá so sánh Baseline & Hoàn thiện báo cáo", 
         "• Cả nhóm: Chạy so sánh Multi-Agent với Single-Agent LLM và Wazuh Rule-based truyền thống\n• Nam: Viết Chương 4 (Triển khai hệ thống EVE-NG, Wazuh & Agent)\n• Đô: Viết Chương 5 (Kết quả thực nghiệm & Đánh giá)", 
         "Bản thảo toàn văn Báo cáo (Chương 1 đến 5)"),
        
        ("13", "13/12 - 14/12", "Chuẩn bị bảo vệ & Đóng gói sản phẩm", 
         "• Nam: Dựng video demo kịch bản tấn công EVE-NG, thiết kế Slide báo cáo\n• Đô: Kiểm tra mã nguồn, in ấn quyển báo cáo, tổng duyệt thuyết trình", 
         "Quyển báo cáo in nộp + Slide + Demo EVE-NG hoàn chỉnh")
    ]

    tbl = doc.add_table(rows=len(table_data) + 1, cols=5)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    # Header
    hdr_row = tbl.rows[0]
    for idx, name in enumerate(headers):
        c = hdr_row.cells[idx]
        c.width = widths[idx]
        set_cell_background(c, "1B365D")
        set_cell_margins(c, 100, 100, 100, 100)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(name)
        r.font.name = 'Times New Roman'
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)

    # Rows
    for r_idx, row_vals in enumerate(table_data):
        row = tbl.rows[r_idx + 1]
        bg = "FFFFFF" if r_idx % 2 == 0 else "F8FAFC"
        for c_idx, val in enumerate(row_vals):
            c = row.cells[c_idx]
            c.width = widths[c_idx]
            set_cell_background(c, bg)
            set_cell_margins(c, 80, 80, 90, 90)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(9.5)
            r.font.color.rgb = COLOR_DARK

    set_table_borders(tbl, "CBD5E1", "4")
    
    out_path = "/home/nam/Desktop/TLCN-NAM/BC/CongViecChinh.docx"
    doc.save(out_path)
    print(f"Updated successfully: {out_path}")

if __name__ == "__main__":
    update_cong_viec_chinh()
